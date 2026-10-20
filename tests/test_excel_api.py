from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook


def test_project_excel_export_download(client):
    project = client.post("/api/projects", json={"nom": "Export web", "takt_s": 55}).json()
    station = client.post(
        f"/api/projects/{project['id']}/stations",
        json={"nom": "Montage"},
    ).json()
    phase = client.post(
        f"/api/stations/{station['id']}/phases",
        json={
            "ordre": 1,
            "debut_s": 2.0,
            "fin_s": 7.0,
            "type_auto": "activity",
            "description": "Assembler",
            "categorie": "VA",
            "chrono_reel_s": 4.0,
        },
    )
    assert phase.status_code == 201, phase.text
    cycle = client.post(
        f"/api/stations/{station['id']}/cycles",
        json={"debut_s": 2.0, "fin_s": 7.0},
    )
    assert cycle.status_code == 201, cycle.text

    response = client.get(f"/api/projects/{project['id']}/export.xlsx")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"] == (
        "attachment; filename*=utf-8''kaizen_studio_Export%20web.xlsx"
    )

    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == ["Synthèse", "Étapes", "Notes"]
    assert workbook["Étapes"]["E2"].value == "=D2-C2"
    assert workbook["Synthèse"]["B6"].value.startswith("=SUMIFS(")
    assert workbook["Notes"]["A11"].value == "Poste"
    assert workbook["Notes"]["E12"].value == "=D12-C12"
    assert workbook["Notes"]["A12"].value == "Montage"


def test_project_excel_export_returns_not_found(client):
    response = client.get("/api/projects/99999/export.xlsx")
    assert response.status_code == 404
