from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "presentation" / "Kaizen_Studio_Direction_v3_Pilote.pptx"

NAVY = "11263D"
NAVY_LIGHT = "1D3B57"
TEAL = "19A890"
TEAL_LIGHT = "DDF5EF"
BLUE = "4C8DCE"
BLUE_LIGHT = "E8F1FA"
AMBER = "F0B44D"
AMBER_LIGHT = "FFF3D8"
RED = "D76C5B"
RED_LIGHT = "FCE9E5"
INK = "1D2B3A"
MUTED = "647487"
PAPER = "F4F7FA"
WHITE = "FFFFFF"
LINE = "DCE4EB"
GREY = "9BA8B5"
FONT = "Aptos"


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_background(slide, value: str = PAPER) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color(value)


def textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    size: float = 16,
    fill: str = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0,
) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(fill)
    return shape


def rounded_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: str,
    line: str | None = None,
    radius_shape=MSO_SHAPE.ROUNDED_RECTANGLE,
):
    shape = slide.shapes.add_shape(
        radius_shape,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.color.rgb = color(line or fill)
    return shape


def line(slide, x1: float, y1: float, x2: float, y2: float, stroke: str, width: float = 1.5) -> None:
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    shape.line.color.rgb = color(stroke)
    shape.line.width = Pt(width)


def add_header(slide, section: str, title: str, subtitle: str | None = None) -> None:
    textbox(slide, 0.62, 0.35, 9.4, 0.24, section.upper(), 9, TEAL, True)
    textbox(slide, 0.62, 0.7, 12.0, 0.52, title, 25, NAVY, True)
    if subtitle:
        textbox(slide, 0.64, 1.28, 12.0, 0.34, subtitle, 11, MUTED)
    line(slide, 0.62, 1.77, 12.72, 1.77, LINE, 1)


def add_footer(slide, number: int) -> None:
    line(slide, 0.62, 7.08, 12.72, 7.08, LINE, 0.8)
    textbox(slide, 0.64, 7.13, 9.5, 0.2, "KAIZEN STUDIO  ·  PROJET PILOTE  ·  OCTOBRE 2026", 8, MUTED, True)
    textbox(slide, 12.0, 7.12, 0.68, 0.22, f"{number:02d} / 10", 8, MUTED, True, PP_ALIGN.RIGHT)


def new_content_slide(prs: Presentation, number: int, section: str, title: str, subtitle: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, section, title, subtitle)
    add_footer(slide, number)
    return slide


def card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    title: str,
    body: str,
    accent: str = TEAL,
    title_size: float = 15,
    body_size: float = 11,
) -> None:
    rounded_rect(slide, x, y, w, h, WHITE, LINE)
    rounded_rect(slide, x, y, 0.07, h, accent, accent, MSO_SHAPE.RECTANGLE)
    textbox(slide, x + 0.22, y + 0.2, w - 0.42, 0.42, title, title_size, NAVY, True)
    textbox(slide, x + 0.22, y + 0.72, w - 0.42, h - 0.88, body, body_size, MUTED)


def bullet(slide, x: float, y: float, text: str, width: float = 5.5, size: float = 13, dot_color: str = TEAL) -> None:
    rounded_rect(slide, x, y + 0.12, 0.09, 0.09, dot_color, dot_color, MSO_SHAPE.OVAL)
    textbox(slide, x + 0.23, y, width - 0.23, 0.45, text, size, INK)


def add_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)
    rounded_rect(slide, 9.18, 0, 4.153, 7.5, NAVY_LIGHT, NAVY_LIGHT)
    rounded_rect(slide, 10.05, 1.15, 2.45, 2.45, TEAL, TEAL)
    textbox(slide, 10.05, 1.55, 2.45, 1.1, "KS", 46, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    textbox(slide, 0.85, 0.78, 5.7, 0.3, "PROJET D’AMÉLIORATION INDUSTRIELLE", 10, "8FE2D2", True)
    textbox(slide, 0.85, 1.65, 7.9, 1.0, "Kaizen Studio", 34, WHITE, True)
    textbox(
        slide,
        0.9,
        2.82,
        7.1,
        1.15,
        "De la vidéo aux décisions d’amélioration",
        22,
        WHITE,
    )
    textbox(
        slide,
        0.92,
        4.3,
        6.6,
        0.75,
        "Remplacer le visionnage avec pauses et la ressaisie manuelle par un flux de revue structuré, traçable et partageable.",
        13,
        "C7D4DF",
    )
    rounded_rect(slide, 0.9, 5.55, 3.25, 0.48, NAVY_LIGHT, NAVY_LIGHT)
    textbox(slide, 1.08, 5.67, 2.9, 0.18, "PRÉSENTATION À LA DIRECTION", 9, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, 0.92, 6.63, 5.2, 0.25, "Octobre 2026  ·  Proposition de pilote", 10, "AAB9C7")
    line(slide, 10.0, 4.15, 12.55, 4.15, "52718A", 1.2)
    textbox(slide, 10.04, 4.45, 2.5, 0.76, "MESURER\n·  COMPRENDRE\n·  AMÉLIORER", 11, WHITE, True)
    textbox(slide, 0.92, 7.14, 10.5, 0.18, "Les temps issus de la vidéo sont des estimations à valider par chronométrage.", 8, "AAB9C7")


def slide_problem(prs: Presentation) -> None:
    slide = new_content_slide(
        prs,
        2,
        "01 · Enjeu",
        "Aujourd’hui : une analyse utile, mais fragmentée",
        "La pratique actuelle — vidéo, pauses et feuille Excel — fonctionne, mais demande beaucoup d’allers-retours manuels.",
    )
    rounded_rect(slide, 0.78, 2.12, 5.7, 3.78, WHITE, LINE)
    textbox(slide, 1.08, 2.42, 4.9, 0.35, "LA MÉTHODE ACTUELLE", 13, NAVY, True)
    current_steps = [
        ("1", "Lire la vidéo", "Repérer une séquence à mesurer."),
        ("2", "Mettre en pause", "Noter le temps dans Excel."),
        ("3", "Reprendre / reculer", "Rechercher le prochain changement."),
        ("4", "Ressaisir et calculer", "Assembler ensuite les données et les totaux."),
    ]
    for index, (number, title, body) in enumerate(current_steps):
        y = 3.05 + index * 0.62
        rounded_rect(slide, 1.1, y, 0.36, 0.36, BLUE_LIGHT, BLUE_LIGHT, MSO_SHAPE.OVAL)
        textbox(slide, 1.1, y + 0.1, 0.36, 0.14, number, 8, BLUE, True, PP_ALIGN.CENTER)
        textbox(slide, 1.65, y - 0.01, 2.0, 0.24, title, 11, NAVY, True)
        textbox(slide, 3.65, y - 0.01, 2.4, 0.34, body, 9, MUTED)
    rounded_rect(slide, 6.82, 2.12, 5.7, 3.78, NAVY, NAVY)
    textbox(slide, 7.15, 2.42, 4.95, 0.35, "CE QUE CELA IMPLIQUE", 13, "8FE2D2", True)
    implications = [
        ("01", "Changements de contexte", "L’attention alterne entre vidéo et tableur."),
        ("02", "Temps d’analyse difficile à suivre", "Les reprises et ressaisies varient selon la méthode."),
        ("03", "Comparaison à reconstruire", "Le takt et les postes demandent une mise en forme."),
        ("04", "Capitalisation dépendante du fichier", "La revue et les corrections sont moins centralisées."),
    ]
    for index, (number, title, body) in enumerate(implications):
        y = 3.02 + index * 0.65
        textbox(slide, 7.17, y, 0.43, 0.2, number, 9, "8FE2D2", True)
        textbox(slide, 7.75, y - 0.02, 4.15, 0.22, title, 10, WHITE, True)
        textbox(slide, 7.75, y + 0.24, 4.15, 0.3, body, 8, "C7D4DF")
    textbox(slide, 0.84, 6.2, 11.5, 0.42, "Ces irritants sont des hypothèses à vérifier avec les utilisateurs ; aucun gain chiffré n’est revendiqué à ce stade.", 10, MUTED)


def slide_workflow(prs: Presentation) -> None:
    slide = new_content_slide(
        prs,
        3,
        "02 · Fonctionnement",
        "Demain : une étude dans un seul flux",
        "Kaizen Studio rassemble la vidéo, les phases, la qualification des temps et la synthèse — l’ingénieur garde la décision.",
    )
    stages = [
        ("01", "Importer", "Vidéo locale\npar poste", BLUE),
        ("02", "Détecter", "Mouvement et\narrêts estimés", TEAL),
        ("03", "Revoir", "Corriger, décrire,\nclasser et chronométrer", AMBER),
        ("04", "Comparer", "Synthèse poste,\ncycle et takt", BLUE),
        ("05", "Partager", "Yamazumi et\nexport Excel", TEAL),
    ]
    start_x, width, gap, y = 0.7, 2.18, 0.32, 2.7
    for index, (number, title, body, accent) in enumerate(stages):
        x = start_x + index * (width + gap)
        rounded_rect(slide, x, y, width, 1.9, WHITE, LINE)
        rounded_rect(slide, x + 0.2, y + 0.2, 0.45, 0.45, accent, accent, MSO_SHAPE.OVAL)
        textbox(slide, x + 0.2, y + 0.31, 0.45, 0.16, number, 8, WHITE, True, PP_ALIGN.CENTER)
        textbox(slide, x + 0.2, y + 0.82, width - 0.4, 0.31, title, 15, NAVY, True)
        textbox(slide, x + 0.2, y + 1.23, width - 0.4, 0.56, body, 10, MUTED)
        if index < len(stages) - 1:
            line(slide, x + width + 0.03, y + 0.95, x + width + gap - 0.04, y + 0.95, GREY, 1.6)
    rounded_rect(slide, 1.0, 5.28, 11.3, 0.72, TEAL_LIGHT, TEAL_LIGHT)
    textbox(slide, 1.22, 5.48, 10.9, 0.25, "L’automate repère seulement la présence de mouvement. L’analyste nomme, classe, corrige et valide chaque opération.", 11, NAVY, True)
    textbox(slide, 0.75, 6.26, 11.4, 0.34, "Résultat attendu : moins d’allers-retours et une analyse plus directement exploitable — à confirmer par le pilote.", 10, MUTED)


def slide_mvp(prs: Presentation) -> None:
    slide = new_content_slide(
        prs,
        4,
        "03 · MVP disponible",
        "Une base déjà testable, une reconnaissance encore expérimentale",
        "Distinguer clairement les fonctions opérationnelles des capacités à évaluer sur les vidéos de l’usine.",
    )
    items = [
        ("Projets & postes", "Takt par projet, vidéo par poste, durée et cadence lues à l’import.", BLUE),
        ("Repérage vidéo", "OpenCV repère activité / arrêt ; seuil et durée minimale sont réglables.", TEAL),
        ("Revue humaine", "Frise synchronisée, coupe/fusion, ajustement des limites, catégories, commentaires et chronos réels.", AMBER),
        ("Décision & partage", "Synthèse VA/NVA/gaspillage, cycles, Yamazumi et export Excel avec formules.", BLUE),
        ("Posture en appui", "Squelette 2D et angles estimés avec MediaPipe ; cotations RULA/REBA saisies par l’évaluateur.", TEAL),
        ("Apprentissage local", "Classifieur expérimental à entraîner sur annotations ; propositions à accepter ou corriger.", AMBER),
    ]
    positions = [(0.72, 2.13), (4.78, 2.13), (8.84, 2.13), (0.72, 4.22), (4.78, 4.22), (8.84, 4.22)]
    for (title, body, accent), (x, y) in zip(items, positions):
        card(slide, x, y, 3.78, 1.72, title, body, accent, 14, 10)


def slide_yamazumi(prs: Presentation) -> None:
    slide = new_content_slide(
        prs,
        5,
        "04 · Yamazumi",
        "Du relevé à une lecture d’équilibrage",
        "Le Yamazumi rend comparables, dans une vue, la composition des temps et l’écart au takt.",
    )
    chart_x, chart_y, chart_w, chart_h = 1.08, 2.38, 7.55, 3.55
    base_y = chart_y + chart_h
    max_value = 90
    scale = chart_h / max_value
    for tick in (0, 20, 40, 60, 80):
        y = base_y - tick * scale
        line(slide, chart_x + 0.46, y, chart_x + chart_w, y, LINE, 0.7)
        textbox(slide, chart_x - 0.08, y - 0.11, 0.42, 0.2, str(tick), 8, MUTED, False, PP_ALIGN.RIGHT)
    takt_y = base_y - 60 * scale
    line(slide, chart_x + 0.48, takt_y, chart_x + chart_w, takt_y, RED, 2)
    textbox(slide, chart_x + chart_w - 0.85, takt_y - 0.27, 0.82, 0.2, "TAKT 60 s", 8, RED, True, PP_ALIGN.RIGHT)

    example = [
        ("Poste A", [("VA", 38, TEAL), ("NVA", 12, BLUE), ("Gaspillage", 18, AMBER)]),
        ("Poste B", [("VA", 44, TEAL), ("NVA", 10, BLUE), ("Gaspillage", 13, AMBER)]),
        ("Poste C", [("VA", 28, TEAL), ("NVA", 22, BLUE), ("Gaspillage", 20, AMBER)]),
    ]
    bar_w = 0.95
    for index, (station, segments) in enumerate(example):
        x = chart_x + 1.1 + index * 2.0
        cursor = base_y
        for label, value, fill in segments:
            segment_h = value * scale
            cursor -= segment_h
            rounded_rect(slide, x, cursor, bar_w, segment_h, fill, fill, MSO_SHAPE.RECTANGLE)
            textbox(slide, x, cursor + max(0, segment_h / 2 - 0.09), bar_w, 0.18, f"{value}", 9, WHITE, True, PP_ALIGN.CENTER)
        textbox(slide, x - 0.2, base_y + 0.13, bar_w + 0.4, 0.22, station, 9, NAVY, True, PP_ALIGN.CENTER)
    legend = [("VA", TEAL), ("NVA nécessaire", BLUE), ("Gaspillage", AMBER)]
    for index, (label, fill) in enumerate(legend):
        x = chart_x + 0.65 + index * 1.65
        rounded_rect(slide, x, 6.38, 0.16, 0.16, fill, fill, MSO_SHAPE.RECTANGLE)
        textbox(slide, x + 0.24, 6.35, 1.42, 0.2, label, 8, MUTED)

    rounded_rect(slide, 9.15, 2.45, 3.25, 2.8, WHITE, LINE)
    textbox(slide, 9.45, 2.75, 2.65, 0.48, "Pourquoi c’est utile", 14, NAVY, True)
    bullet(slide, 9.45, 3.43, "Visualiser un poste au-dessus du takt.", 2.75, 10, RED)
    bullet(slide, 9.45, 4.08, "Distinguer temps VA, NVA et gaspillage.", 2.75, 10, TEAL)
    bullet(slide, 9.45, 4.73, "Orienter la discussion vers l’équilibrage.", 2.75, 10, BLUE)
    rounded_rect(slide, 9.15, 5.48, 3.25, 0.73, AMBER_LIGHT, AMBER_LIGHT)
    textbox(slide, 9.38, 5.67, 2.8, 0.34, "Exemple fictif — non issu d’une étude usine.", 9, NAVY, True)
    textbox(slide, 0.95, 6.74, 8.0, 0.18, "Les durées restent des estimations vidéo à vérifier et à compléter par chronométrage réel.", 8, MUTED)


def slide_data_quality(prs: Presentation) -> None:
    slide = new_content_slide(
        prs,
        6,
        "05 · Qualité de l’analyse",
        "Une aide à l’analyse, pas une boîte noire",
        "Les temps vidéo restent des estimations ; l’application organise la revue et la validation humaine.",
    )
    rounded_rect(slide, 0.78, 2.22, 5.55, 3.52, BLUE_LIGHT, BLUE_LIGHT)
    textbox(slide, 1.12, 2.57, 4.7, 0.42, "ESTIMATION VIDÉO", 16, NAVY, True)
    bullet(slide, 1.15, 3.28, "Signal de mouvement et phases activité / arrêt.", 4.7, 12, BLUE)
    bullet(slide, 1.15, 3.92, "Sensibles au cadrage, à la lumière et aux occultations.", 4.7, 12, BLUE)
    bullet(slide, 1.15, 4.56, "Modèle d’opérations expérimental, à évaluer.", 4.7, 12, BLUE)
    rounded_rect(slide, 6.85, 2.22, 5.55, 3.52, TEAL_LIGHT, TEAL_LIGHT)
    textbox(slide, 7.2, 2.57, 4.7, 0.42, "VALIDATION HUMAINE", 16, NAVY, True)
    bullet(slide, 7.24, 3.28, "L’analyste nomme et classe les opérations.", 4.7, 12, TEAL)
    bullet(slide, 7.24, 3.92, "Chronométrage réel saisi pour comparer les écarts.", 4.7, 12, TEAL)
    bullet(slide, 7.24, 4.56, "Suggestions du modèle vérifiées avant acceptation.", 4.7, 12, TEAL)
    rounded_rect(slide, 1.2, 6.06, 10.95, 0.55, NAVY, NAVY)
    textbox(slide, 1.45, 6.22, 10.45, 0.2, "Principe : faciliter la revue, rendre les corrections visibles, valider au chronomètre.", 12, WHITE, True, PP_ALIGN.CENTER)


def slide_security(prs: Presentation) -> None:
    slide = new_content_slide(
        prs,
        7,
        "06 · Données & déploiement",
        "La confidentialité fait partie de la proposition de valeur",
        "Les vidéos d’atelier sont sensibles : garder la maîtrise des données est un prérequis au pilote.",
    )
    card(slide, 0.78, 2.2, 3.75, 2.15, "Local par défaut", "Serveur sur le PC ; vidéos, base et modèles stockés localement. Aucun appel externe requis pour le flux courant.", TEAL, 15, 11)
    card(slide, 4.78, 2.2, 3.75, 2.15, "Accès maîtrisé", "Comptes locaux, rôles et séparation des projets par utilisateur. Usage multi-PC à cadrer avec l’IT.", BLUE, 15, 11)
    card(slide, 8.78, 2.2, 3.75, 2.15, "Cadre social & RGPD", "Informer les opérateurs ; définir l’accès, la durée de conservation et les conditions d’utilisation.", AMBER, 15, 11)
    rounded_rect(slide, 0.82, 4.78, 11.7, 1.18, RED_LIGHT, RED_LIGHT)
    textbox(slide, 1.12, 5.03, 11.1, 0.3, "Avant tout déploiement élargi", 13, NAVY, True)
    textbox(slide, 1.12, 5.43, 11.0, 0.36, "Faire valider l’architecture, les droits d’accès, la conservation des vidéos et l’information des personnes par les interlocuteurs compétents.", 11, INK)
    textbox(slide, 0.86, 6.4, 11.5, 0.25, "Le stockage local réduit les transferts, mais ne remplace ni les règles internes ni l’analyse RGPD.", 9, MUTED)


def slide_roadmap(prs: Presentation) -> None:
    slide = new_content_slide(
        prs,
        8,
        "07 · Pilote terrain",
        "Prouver la valeur par une comparaison simple et équitable",
        "Mesurer le travail réel avec et sans Kaizen Studio, sur le même périmètre et selon les mêmes règles.",
    )
    phases = [
        ("1", "Choisir un cas représentatif", "Un poste, une vidéo et un cycle suffisamment compris par l’équipe.", BLUE),
        ("2", "Établir la référence", "Réaliser l’étude selon la pratique vidéo + pauses + Excel actuelle.", TEAL),
        ("3", "Rejouer avec Kaizen Studio", "Même périmètre, mêmes conventions et analyste comparable.", AMBER),
        ("4", "Comparer la qualité", "Vérifier limites, catégories, chronos et corrections avec un référent méthodes.", BLUE),
        ("5", "Décider sur les faits", "Comparer temps d’analyse, complétude, reprises et utilité de la synthèse.", TEAL),
    ]
    row_y = 1.98
    row_height = 0.72
    row_gap = 0.09
    for index, (number, title, body, accent) in enumerate(phases):
        y = row_y + index * (row_height + row_gap)
        rounded_rect(slide, 0.78, y, 11.75, row_height, WHITE, LINE)
        rounded_rect(slide, 0.96, y + 0.13, 0.44, 0.44, accent, accent, MSO_SHAPE.OVAL)
        textbox(slide, 0.96, y + 0.25, 0.44, 0.15, number, 9, WHITE, True, PP_ALIGN.CENTER)
        textbox(slide, 1.6, y + 0.12, 2.55, 0.43, title, 12, NAVY, True, valign=MSO_ANCHOR.MIDDLE)
        textbox(slide, 4.15, y + 0.11, 8.05, 0.48, body, 10, MUTED, valign=MSO_ANCHOR.MIDDLE)
    rounded_rect(slide, 0.82, 6.08, 11.7, 0.52, AMBER_LIGHT, AMBER_LIGHT)
    textbox(slide, 1.05, 6.22, 11.2, 0.23, "Le pilote valide l’outil et la méthode ; il ne présuppose ni gain de temps ni gain de productivité.", 9, NAVY, True, PP_ALIGN.CENTER)


def slide_pilot(prs: Presentation) -> None:
    slide = new_content_slide(
        prs,
        9,
        "08 · Indicateurs de réussite",
        "Ce que la direction pourra objectivement évaluer",
        "Pas de promesse au départ : une mesure avant / après transforme une intuition en décision d’investissement.",
    )
    steps = [
        ("01", "Temps d’étude", "Durée totale analyste, du chargement à l’export.", BLUE),
        ("02", "Qualité des limites", "Écart des débuts/fins aux chronos de référence.", TEAL),
        ("03", "Fiabilité des données", "Étapes manquantes, doublons, corrections requises.", AMBER),
        ("04", "Utilité décisionnelle", "Clarté du Yamazumi et temps nécessaire pour préparer la synthèse.", BLUE),
    ]
    for index, (number, title, body, accent) in enumerate(steps):
        x = 0.8 + index * 3.12
        rounded_rect(slide, x, 2.46, 2.77, 2.0, WHITE, LINE)
        textbox(slide, x + 0.23, 2.72, 0.48, 0.25, number, 11, accent, True)
        textbox(slide, x + 0.23, 3.16, 2.25, 0.4, title, 15, NAVY, True)
        textbox(slide, x + 0.23, 3.68, 2.25, 0.54, body, 10, MUTED)
        if index < 3:
            line(slide, x + 2.8, 3.45, x + 3.05, 3.45, GREY, 1.5)
    rounded_rect(slide, 0.82, 4.92, 11.7, 1.22, NAVY, NAVY)
    textbox(slide, 1.12, 5.15, 2.05, 0.3, "Protocole", 12, "8FE2D2", True)
    textbox(slide, 3.25, 5.08, 8.85, 0.7, "Même vidéo · mêmes règles de découpage · analyste formé · chronométrage des deux méthodes · contrôle des résultats par le référent méthodes", 11, WHITE, True, valign=MSO_ANCHOR.MIDDLE)
    textbox(slide, 0.86, 6.4, 11.5, 0.24, "Définir avant l’essai les critères d’acceptation avec méthodes, production et direction.", 9, MUTED)


def slide_decision(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)
    textbox(slide, 0.8, 0.55, 11.8, 0.27, "09 · DÉCISION ATTENDUE", 10, "8FE2D2", True)
    textbox(slide, 0.8, 1.06, 11.8, 0.72, "Autoriser un pilote mesuré", 30, WHITE, True)
    textbox(slide, 0.84, 1.96, 11.3, 0.52, "Tester Kaizen Studio sur un cas réel avant de décider d’un déploiement plus large.", 15, "C7D4DF")

    decisions = [
        ("01", "Périmètre", "Quel poste, quel cycle et quel référent méthodes ?"),
        ("02", "Ressources", "Qui réalise les deux études et valide les temps de référence ?"),
        ("03", "Garde-fous", "Quelles règles vidéo, durée de conservation et information opérateurs ?"),
    ]
    for index, (number, title, body) in enumerate(decisions):
        y = 3.0 + index * 0.9
        rounded_rect(slide, 0.86, y, 0.52, 0.52, TEAL, TEAL, MSO_SHAPE.OVAL)
        textbox(slide, 0.86, y + 0.16, 0.52, 0.18, number, 9, WHITE, True, PP_ALIGN.CENTER)
        textbox(slide, 1.63, y + 0.02, 2.0, 0.3, title, 14, WHITE, True)
        textbox(slide, 3.55, y + 0.02, 7.9, 0.44, body, 12, "C7D4DF")

    rounded_rect(slide, 0.86, 5.98, 11.5, 0.65, NAVY_LIGHT, NAVY_LIGHT)
    textbox(slide, 1.15, 6.17, 10.9, 0.25, "Décision proposée : désigner un référent et sélectionner une vidéo représentative pour la comparaison.", 11, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, 0.9, 7.13, 8.0, 0.18, "Kaizen Studio · Les temps automatiques restent des estimations à valider.", 8, "AAB9C7")
    textbox(slide, 12.0, 7.12, 0.68, 0.22, "10 / 10", 8, "AAB9C7", True, PP_ALIGN.RIGHT)


def build_presentation() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Kaizen Studio — Proposition de pilote à la direction"
    prs.core_properties.subject = "Du visionnage manuel et Excel à une étude de temps structurée et mesurable"
    prs.core_properties.author = "Kaizen Studio"

    add_title_slide(prs)
    slide_problem(prs)
    slide_workflow(prs)
    slide_mvp(prs)
    slide_yamazumi(prs)
    slide_data_quality(prs)
    slide_security(prs)
    slide_roadmap(prs)
    slide_pilot(prs)
    slide_decision(prs)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Présentation créée : {OUTPUT}")


if __name__ == "__main__":
    build_presentation()
