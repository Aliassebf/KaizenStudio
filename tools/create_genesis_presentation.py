from __future__ import annotations

import os
import sqlite3
import tempfile
from collections import defaultdict
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path.home() / "OneDrive - Otis Global Tenant" / "Otis industriel" / "Gien template 2026.potx"
OUTPUT = ROOT / "presentation" / "Kaizen_Studio_GENESIS_Direction_Otis_20261015.pptx"

NAVY = "041E42"
BLUE = "0033A0"
GOLD = "CBA052"
GREY = "6E747B"
LIGHT_GREY = "F3F4F4"
LINE = "D8DADD"
WHITE = "FFFFFF"
INK = "252A30"
GREEN = "37846A"
LIGHT_GREEN = "E7F2ED"
LIGHT_GOLD = "F7F0E3"
BODY_FONT = "Century Gothic"
TITLE_FONT = "Georgia"
BLANK_LAYOUT = 28
COVER_LAYOUT = 2


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def textbox(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    text: str,
    size: float = 12,
    color: str = INK,
    bold: bool = False,
    font: str = BODY_FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def rect(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    fill: str = WHITE,
    outline: str = LINE,
    shape_type=MSO_SHAPE.ROUNDED_RECTANGLE,
):
    shape = slide.shapes.add_shape(
        shape_type,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(outline)
    return shape


def rule(slide, x1: float, y1: float, x2: float, y2: float, color: str = LINE, width: float = 1):
    shape = slide.shapes.add_connector(
        1,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)


def load_genesis_data() -> dict:
    data_dir = os.environ.get("KAIZEN_STUDIO_DATA_DIR")
    if data_dir:
        database = Path(data_dir).expanduser() / "data" / "kaizen_studio.db"
    else:
        local_app_data = os.environ.get("LOCALAPPDATA")
        if not local_app_data:
            raise RuntimeError("Le dossier de données local de Kaizen Studio est introuvable.")
        database = Path(local_app_data) / "Kaizen Studio" / "data" / "kaizen_studio.db"
    if not database.is_file():
        raise FileNotFoundError("La base locale Kaizen Studio est absente ; impossible de lire GENESIS.")

    connection = sqlite3.connect(f"{database.as_uri()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    try:
        project = connection.execute(
            "SELECT id, nom, takt_s FROM projects WHERE nom = ? COLLATE NOCASE ORDER BY id DESC LIMIT 1",
            ("GENESIS",),
        ).fetchone()
        if project is None:
            raise LookupError("Le projet GENESIS n’existe pas dans la base locale.")

        stations = connection.execute(
            "SELECT id, nom, video_duration_s, video_fps FROM stations WHERE project_id = ? ORDER BY id",
            (project["id"],),
        ).fetchall()
        station_ids = [station["id"] for station in stations]
        phases = []
        cycles = 0
        measured_phases = 0
        if station_ids:
            placeholders = ",".join("?" for _ in station_ids)
            phases = connection.execute(
                f"SELECT ordre, debut_s, fin_s, description, categorie, chrono_reel_s "
                f"FROM phases WHERE station_id IN ({placeholders}) ORDER BY ordre",
                station_ids,
            ).fetchall()
            cycles = connection.execute(
                f"SELECT COUNT(*) FROM cycles WHERE station_id IN ({placeholders})",
                station_ids,
            ).fetchone()[0]
            measured_phases = connection.execute(
                f"SELECT COUNT(*) FROM phases WHERE station_id IN ({placeholders}) "
                "AND chrono_reel_s IS NOT NULL",
                station_ids,
            ).fetchone()[0]

        category_seconds: dict[str, float] = defaultdict(float)
        phase_summary = []
        for phase in phases:
            duration = max(0.0, float(phase["fin_s"]) - float(phase["debut_s"]))
            category_seconds[phase["categorie"] or "Non classé"] += duration
            phase_summary.append({
                "ordre": int(phase["ordre"]),
                "description": (phase["description"] or "Étape non décrite").strip(),
                "categorie": phase["categorie"] or "Non classé",
                "duree_s": duration,
            })

        return {
            "project": dict(project),
            "stations": [dict(station) for station in stations],
            "phases": phase_summary,
            "category_seconds": dict(category_seconds),
            "video_duration_s": max(
                (float(station["video_duration_s"] or 0) for station in stations),
                default=0,
            ),
            "cycles": int(cycles),
            "measured_phases": int(measured_phases),
        }
    finally:
        connection.close()


def open_template() -> Presentation:
    if not TEMPLATE.is_file():
        raise FileNotFoundError(f"Le template PowerPoint demandé est introuvable : {TEMPLATE}")
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as temporary:
            temporary_path = Path(temporary.name)
        with ZipFile(TEMPLATE) as source, ZipFile(temporary_path, "w", ZIP_DEFLATED) as destination:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                        b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                    )
                destination.writestr(item, data)
        presentation = Presentation(str(temporary_path))
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)

    slide_ids = presentation.slides._sldIdLst
    for slide_id in list(slide_ids):
        presentation.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)
    return presentation


def add_header(slide, section: str, title: str, subtitle: str = "") -> None:
    textbox(slide, 0.52, 0.22, 8.95, 0.2, section.upper(), 8, BLUE, True)
    textbox(slide, 0.52, 0.5, 8.95, 0.49, title, 21, NAVY, True, TITLE_FONT)
    if subtitle:
        textbox(slide, 0.54, 1.04, 8.85, 0.34, subtitle, 9.5, GREY)
    rule(slide, 0.52, 1.48, 9.45, 1.48, GOLD, 1.4)


def content_slide(presentation: Presentation, section: str, title: str, subtitle: str = ""):
    slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK_LAYOUT])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("FAFAF8")
    add_header(slide, section, title, subtitle)
    return slide


def card(slide, x, y, width, height, title, body, accent=GOLD, fill=WHITE):
    rect(slide, x, y, width, height, fill, LINE)
    rect(slide, x, y, 0.07, height, accent, accent, MSO_SHAPE.RECTANGLE)
    textbox(slide, x + 0.2, y + 0.16, width - 0.38, 0.36, title, 12, NAVY, True)
    textbox(slide, x + 0.2, y + 0.56, width - 0.38, height - 0.68, body, 9.4, INK)


def metric_card(slide, x, y, width, value, label, note="", accent=BLUE):
    rect(slide, x, y, width, 1.05, WHITE, LINE)
    textbox(slide, x + 0.15, y + 0.13, width - 0.3, 0.4, value, 23, accent, True, TITLE_FONT)
    textbox(slide, x + 0.15, y + 0.55, width - 0.3, 0.23, label, 9, NAVY, True)
    if note:
        textbox(slide, x + 0.15, y + 0.8, width - 0.3, 0.18, note, 7.5, GREY)


def add_cover(presentation: Presentation, data: dict) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[COVER_LAYOUT])
    slide.placeholders[11].text = "PROPOSITION DE PILOTE · OCTOBRE 2026"
    slide.placeholders[0].text = "Kaizen Studio"
    slide.placeholders[10].text = (
        "GENESIS · Gien\n"
        "De la vidéo et des pauses Excel à une étude structurée, vérifiable et exploitable."
    )


def slide_problem(presentation: Presentation) -> None:
    slide = content_slide(
        presentation,
        "01 · Constat",
        "Aujourd’hui : vidéo, pauses et ressaisie",
        "Une méthode familière, mais des informations dispersées entre le lecteur vidéo et Excel.",
    )
    steps = [
        ("1", "Lire la vidéo", "Rechercher le début et la fin de chaque opération."),
        ("2", "Mettre en pause", "Noter les temps au fur et à mesure dans une feuille Excel."),
        ("3", "Reprendre", "Revenir dans la vidéo et répéter l’opération sur tout le cycle."),
    ]
    for i, (number, title, body) in enumerate(steps):
        x = 0.58 + i * 3.02
        rect(slide, x, 1.85, 2.72, 1.45)
        rect(slide, x + 0.18, 2.05, 0.43, 0.43, GOLD, GOLD, MSO_SHAPE.OVAL)
        textbox(slide, x + 0.18, 2.17, 0.43, 0.16, number, 9, WHITE, True, align=PP_ALIGN.CENTER)
        textbox(slide, x + 0.74, 2.03, 1.8, 0.32, title, 12, NAVY, True)
        textbox(slide, x + 0.74, 2.43, 1.78, 0.59, body, 9, GREY)
    card(
        slide,
        0.58,
        3.65,
        8.82,
        0.9,
        "Ce que le pilote doit vérifier",
        "Temps de manipulation, cohérence des relevés, reprises nécessaires et facilité de synthèse — sans supposer de gain avant mesure.",
        BLUE,
        LIGHT_GREY,
    )


def slide_workflow(presentation: Presentation) -> None:
    slide = content_slide(
        presentation,
        "02 · Réponse proposée",
        "Un flux unique, avec l’analyste aux commandes",
        "Kaizen Studio facilite le découpage et la synthèse ; l’humain garde la décision sur les opérations.",
    )
    items = [
        ("Vidéo", "Importer le poste"),
        ("Signal", "Repérer activité / arrêt"),
        ("Revue", "Corriger les limites"),
        ("Qualification", "Décrire et classer"),
        ("Synthèse", "Comparer et exporter"),
    ]
    for i, (title, body) in enumerate(items):
        x = 0.55 + i * 1.82
        rect(slide, x, 2.05, 1.62, 1.34, WHITE, LINE)
        textbox(slide, x + 0.12, 2.24, 1.38, 0.28, title, 10, BLUE, True, align=PP_ALIGN.CENTER)
        textbox(slide, x + 0.12, 2.65, 1.38, 0.49, body, 8.5, INK, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if i < len(items) - 1:
            rule(slide, x + 1.64, 2.72, x + 1.79, 2.72, GOLD, 2)
    rect(slide, 0.58, 3.78, 8.82, 0.65, LIGHT_GOLD, LIGHT_GOLD)
    textbox(
        slide,
        0.8,
        3.96,
        8.35,
        0.28,
        "Le système mesure le mouvement ; l’analyste décrit l’opération, vérifie les temps et valide les conclusions.",
        9.3,
        NAVY,
        True,
        align=PP_ALIGN.CENTER,
    )


def slide_genesis_facts(presentation: Presentation, data: dict) -> None:
    slide = content_slide(
        presentation,
        "03 · Cas concret",
        "GENESIS : un premier état exploitable",
        "Chiffres tirés du projet local GENESIS ; ils décrivent le relevé existant, pas un résultat de pilote comparatif.",
    )
    project = data["project"]
    metric_card(slide, 0.58, 1.82, 2.03, str(len(data["stations"])), "poste(s)", "dans le projet GENESIS", BLUE)
    metric_card(slide, 2.78, 1.82, 2.03, f"{data['video_duration_s']:.1f} s", "durée vidéo", "métadonnée vidéo", NAVY)
    metric_card(slide, 4.98, 1.82, 2.03, str(len(data["phases"])), "étapes saisies", "dans la revue actuelle", GREEN)
    metric_card(slide, 7.18, 1.82, 2.03, f"{float(project['takt_s']):.0f} s", "takt projet", "cible déclarée", GOLD)

    category_text = " · ".join(
        f"{name} {seconds:.1f} s"
        for name, seconds in sorted(data["category_seconds"].items(), key=lambda item: item[1], reverse=True)
    ) or "Aucune catégorie saisie."
    rect(slide, 0.58, 3.22, 8.82, 0.85, LIGHT_GREY, LIGHT_GREY)
    textbox(slide, 0.8, 3.38, 8.4, 0.25, "Durées des étapes déjà classées", 10.5, NAVY, True)
    textbox(slide, 0.8, 3.72, 8.4, 0.23, category_text, 10, INK)
    caution = (
        f"{data['cycles']} cycle enregistré · {data['measured_phases']} étape(s) avec chrono réel. "
        "Le total des segments vidéo ne doit donc pas être présenté comme un temps de cycle validé."
    )
    textbox(slide, 0.62, 4.3, 8.7, 0.4, caution, 8.5, GREY)


def slide_genesis_steps(presentation: Presentation, data: dict) -> None:
    slide = content_slide(
        presentation,
        "04 · Lecture des données",
        "Où se concentrent les durées enregistrées ?",
        "Extrait des trois segments les plus longs de GENESIS — durées calculées à partir des limites saisies.",
    )
    longest = sorted(data["phases"], key=lambda item: item["duree_s"], reverse=True)[:3]
    max_duration = max((item["duree_s"] for item in longest), default=1.0) or 1.0
    for i, phase in enumerate(longest):
        y = 1.84 + i * 0.88
        label = f"{phase['ordre']}. {phase['description']}"
        textbox(slide, 0.62, y, 3.1, 0.32, label, 9.4, NAVY, True)
        rect(slide, 3.85, y + 0.02, 4.3, 0.25, LIGHT_GREY, LIGHT_GREY)
        rect(slide, 3.85, y + 0.02, 4.3 * phase["duree_s"] / max_duration, 0.25, BLUE, BLUE, MSO_SHAPE.RECTANGLE)
        textbox(slide, 8.3, y - 0.02, 1.0, 0.32, f"{phase['duree_s']:.1f} s", 10, NAVY, True, align=PP_ALIGN.RIGHT)
    rect(slide, 0.58, 4.58, 8.82, 0.55, LIGHT_GOLD, LIGHT_GOLD)
    textbox(
        slide,
        0.78,
        4.72,
        8.4,
        0.24,
        "Ces segments orientent la revue ; ils ne prouvent pas à eux seuls un goulot ni un potentiel de gain.",
        8.6,
        NAVY,
        True,
        align=PP_ALIGN.CENTER,
    )


def slide_takt(presentation: Presentation, data: dict) -> None:
    slide = content_slide(
        presentation,
        "05 · Takt et cycle",
        "Le takt est connu ; le cycle reste à délimiter",
        "Un temps vidéo complet ne se compare pas directement au takt tant que les limites de cycle ne sont pas validées.",
    )
    metric_card(slide, 0.72, 1.92, 3.8, f"{float(data['project']['takt_s']):.0f} s", "takt du projet", "objectif de cadence déclaré", GOLD)
    metric_card(slide, 5.0, 1.92, 3.8, f"{data['video_duration_s']:.1f} s", "vidéo du poste", "séquence enregistrée", BLUE)
    rect(slide, 0.72, 3.38, 8.08, 0.96, LIGHT_GOLD, LIGHT_GOLD)
    textbox(slide, 0.98, 3.56, 7.58, 0.26, "Étape suivante pour GENESIS", 10.5, NAVY, True)
    textbox(
        slide,
        0.98,
        3.9,
        7.58,
        0.28,
        "Définir et enregistrer un cycle représentatif ; puis produire un Yamazumi interprétable face au takt.",
        9.2,
        INK,
    )


def slide_value(presentation: Presentation) -> None:
    slide = content_slide(
        presentation,
        "06 · Valeur attendue",
        "Ce que Kaizen Studio peut rendre plus simple",
        "Des améliorations de méthode à mesurer pendant le pilote, et non des gains déjà démontrés.",
    )
    card(slide, 0.58, 1.82, 4.2, 1.15, "Revoir sans perdre le contexte", "Frise, vidéo, étapes et limites restent liées ; les phases sont ajustables.", BLUE)
    card(slide, 5.0, 1.82, 4.2, 1.15, "Structurer les observations", "Description, catégorie, commentaire et chrono réel au même endroit.", GREEN)
    card(slide, 0.58, 3.22, 4.2, 1.15, "Produire une synthèse partageable", "Yamazumi et export Excel à partir des données revues.", GOLD)
    card(slide, 5.0, 3.22, 4.2, 1.15, "Garder la validation humaine", "Les suggestions et temps vidéo sont vérifiés avant toute conclusion.", NAVY)


def slide_yamazumi(presentation: Presentation) -> None:
    slide = content_slide(
        presentation,
        "07 · Synthèse",
        "Du relevé au Yamazumi",
        "Le graphique aide à voir la répartition et les écarts ; sa valeur dépend de cycles comparables et de données validées.",
    )
    bars = [
        ("Valeur ajoutée", 0.78, BLUE),
        ("NVA nécessaire", 0.48, GOLD),
        ("Gaspillage / à qualifier", 0.32, GREY),
    ]
    for i, (label, width, fill) in enumerate(bars):
        y = 1.95 + i * 0.6
        textbox(slide, 0.7, y, 2.1, 0.28, label, 9.3, NAVY, True)
        rect(slide, 2.95, y + 0.01, 4.4, 0.28, LIGHT_GREY, LIGHT_GREY)
        rect(slide, 2.95, y + 0.01, 4.4 * width, 0.28, fill, fill, MSO_SHAPE.RECTANGLE)
    rule(slide, 7.15, 1.82, 7.15, 3.65, NAVY, 1.6)
    textbox(slide, 7.3, 1.77, 1.65, 0.35, "Repère takt", 9, NAVY, True)
    rect(slide, 0.7, 4.0, 8.55, 0.57, LIGHT_GREY, LIGHT_GREY)
    textbox(
        slide,
        0.9,
        4.15,
        8.15,
        0.25,
        "GENESIS compte actuellement un poste et aucun cycle enregistré : pas encore de comparaison multi-postes.",
        8.8,
        INK,
        align=PP_ALIGN.CENTER,
    )
    textbox(slide, 0.72, 4.66, 8.5, 0.22, "Illustration du principe Yamazumi — barres non issues d’une mesure GENESIS.", 7.6, GREY)


def slide_safeguards(presentation: Presentation) -> None:
    slide = content_slide(
        presentation,
        "08 · Maîtrise des risques",
        "Des limites explicites et des garde-fous",
        "La confiance vient de la transparence des estimations et de la validation par les personnes compétentes.",
    )
    card(slide, 0.58, 1.82, 2.72, 1.38, "Estimation", "Le mode motion détecte une présence de mouvement ; il ne reconnaît pas le travail réalisé.", BLUE)
    card(slide, 3.62, 1.82, 2.72, 1.38, "Validation", "Les limites, catégories et temps sont vérifiés par l’analyste et confirmés au chronomètre.", GREEN)
    card(slide, 6.66, 1.82, 2.72, 1.38, "Ergonomie", "RULA/REBA est une pré-cotation à vérifier ; ce n’est pas un diagnostic médical.", GOLD)
    rect(slide, 0.58, 3.58, 8.82, 0.82, LIGHT_GOLD, LIGHT_GOLD)
    textbox(slide, 0.82, 3.75, 8.3, 0.23, "Vidéos d’opérateurs : conditions d’usage à cadrer avant le pilote", 10.2, NAVY, True)
    textbox(slide, 0.82, 4.08, 8.3, 0.2, "Information des personnes, accès, conservation et validation HSE / RH / DPO selon les règles internes.", 8.8, INK)


def slide_pilot(presentation: Presentation) -> None:
    slide = content_slide(
        presentation,
        "09 · Protocole",
        "Comparer les deux méthodes sur le même cas",
        "Le projet GENESIS fournit un point de départ ; le pilote doit ajouter un cycle défini et une référence chronométrée.",
    )
    steps = [
        ("01", "Fixer le cycle", "Début / fin et règles communes."),
        ("02", "Établir la référence", "Vidéo + pauses + Excel."),
        ("03", "Rejouer", "Même vidéo avec Kaizen Studio."),
        ("04", "Valider", "Temps réels et qualité par un référent."),
    ]
    for i, (n, title, body) in enumerate(steps):
        x = 0.58 + i * 2.22
        rect(slide, x, 1.98, 2.02, 1.3)
        textbox(slide, x + 0.16, 2.13, 0.46, 0.28, n, 10, GOLD, True)
        textbox(slide, x + 0.16, 2.48, 1.72, 0.3, title, 10.2, NAVY, True)
        textbox(slide, x + 0.16, 2.84, 1.72, 0.32, body, 8, GREY)
    rect(slide, 0.58, 3.68, 8.82, 0.76, NAVY, NAVY)
    textbox(
        slide,
        0.82,
        3.87,
        8.35,
        0.32,
        "Indicateurs : temps d’étude · écart des limites · reprises/corrections · complétude · utilité du Yamazumi",
        9,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )


def slide_decision(presentation: Presentation, data: dict) -> None:
    slide = content_slide(
        presentation,
        "10 · Décision attendue",
        "Autoriser un pilote mesuré sur GENESIS",
        "Une décision d’expérimentation, sans engagement de déploiement ni promesse de gain.",
    )
    questions = [
        ("Périmètre", "Quel cycle GENESIS et quelle vidéo de référence ?"),
        ("Référent", "Qui valide le découpage et les chronos réels ?"),
        ("Garde-fous", "Quelles règles d’accès, conservation et information des opérateurs ?"),
    ]
    for i, (title, body) in enumerate(questions):
        y = 1.86 + i * 0.78
        rect(slide, 0.7, y, 8.55, 0.62, WHITE, LINE)
        textbox(slide, 0.9, y + 0.17, 1.55, 0.24, title, 10, BLUE, True)
        textbox(slide, 2.52, y + 0.17, 6.45, 0.24, body, 9.2, INK)
    rect(slide, 0.7, 4.34, 8.55, 0.48, LIGHT_GREEN, LIGHT_GREEN)
    textbox(
        slide,
        0.88,
        4.46,
        8.2,
        0.21,
        f"État de départ : {len(data['stations'])} poste · {len(data['phases'])} étapes saisies · takt {float(data['project']['takt_s']):.0f} s.",
        8.8,
        GREEN,
        True,
        align=PP_ALIGN.CENTER,
    )


def build_presentation() -> None:
    data = load_genesis_data()
    presentation = open_template()
    if len(presentation.slide_layouts) <= BLANK_LAYOUT:
        raise RuntimeError("Le template ne contient pas les mises en page attendues.")
    presentation.core_properties.title = "Kaizen Studio — proposition de pilote GENESIS"
    presentation.core_properties.subject = "Étude de temps vidéo, données GENESIS et protocole de validation"
    presentation.core_properties.author = "Kaizen Studio"

    add_cover(presentation, data)
    slide_problem(presentation)
    slide_workflow(presentation)
    slide_genesis_facts(presentation, data)
    slide_genesis_steps(presentation, data)
    slide_takt(presentation, data)
    slide_value(presentation)
    slide_yamazumi(presentation)
    slide_safeguards(presentation)
    slide_pilot(presentation)
    slide_decision(presentation, data)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    print(f"Présentation créée : {OUTPUT}")
    print(f"Diapositives : {len(presentation.slides)}")


if __name__ == "__main__":
    build_presentation()
