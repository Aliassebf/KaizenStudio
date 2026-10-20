from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


CATEGORIES = ["VA", "NVA nécessaire", "Gaspillage", "Non classé"]
CATEGORY_COLORS = {
    "VA": "2E7D32",
    "NVA nécessaire": "F9A825",
    "Gaspillage": "C62828",
    "Non classé": "9E9E9E",
}
FONT_NAME = "Arial"


def _safe_text(value: Any) -> Any:
    """Empêche qu'un texte saisi commençant par « = » soit écrit comme une formule Excel."""
    if isinstance(value, str) and value.lstrip().startswith("="):
        return "'" + value
    return value
EXTRA_STEP_ROWS = 60


def _style_header(worksheet, row: int, column_count: int) -> None:
    border_side = Side(style="thin", color="BFBFBF")
    for column in range(1, column_count + 1):
        cell = worksheet.cell(row=row, column=column)
        cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=border_side,
            bottom=border_side,
            left=border_side,
            right=border_side,
        )


def build_project_workbook(
    project: dict[str, Any],
    stations: list[dict[str, Any]],
    output_path: str | Path,
) -> Path:
    """Crée le classeur du projet avec formules Excel, tableau Yamazumi et notes."""
    output_path = Path(output_path)
    workbook = Workbook()
    steps_sheet = workbook.active
    steps_sheet.title = "Étapes"
    headers = [
        "Poste",
        "N°",
        "Début (s)",
        "Fin (s)",
        "Durée (s)",
        "Description",
        "Catégorie",
        "Commentaire",
        "Type automatique",
        "Chrono réel (s)",
        "Écart auto vs réel (%)",
    ]
    steps_sheet.append(headers)
    _style_header(steps_sheet, 1, len(headers))

    thin_border = Border(
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
    )
    row_number = 2
    for station in stations:
        for phase in station["phases"]:
            steps_sheet.append(
                [
                    _safe_text(station["nom"]),
                    phase["ordre"],
                    phase["debut_s"],
                    phase["fin_s"],
                    f"=D{row_number}-C{row_number}",
                    _safe_text(phase.get("description") or ""),
                    phase.get("categorie") or "Non classé",
                    _safe_text(phase.get("commentaire") or ""),
                    phase["type_auto"],
                    phase.get("chrono_reel_s"),
                    f'=IF(OR(J{row_number}="",J{row_number}=0),"",ABS((E{row_number}-J{row_number})/J{row_number}))',
                ]
            )
            for column in range(1, len(headers) + 1):
                cell = steps_sheet.cell(row=row_number, column=column)
                cell.font = Font(
                    name=FONT_NAME,
                    size=10,
                    color="0000FF" if column in (1, 2, 3, 4, 6, 7, 8, 9, 10) else "000000",
                )
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=column in (6, 8))
            for column in (3, 4, 5, 10):
                steps_sheet.cell(row=row_number, column=column).number_format = "0.0"
            steps_sheet.cell(row=row_number, column=11).number_format = "0.0%"
            row_number += 1

    last_step_row = row_number - 1
    formula_end_row = max(last_step_row + EXTRA_STEP_ROWS, 1 + EXTRA_STEP_ROWS)
    for row in range(2, formula_end_row + 1):
        if steps_sheet.cell(row=row, column=5).value is None:
            steps_sheet.cell(row=row, column=5, value=f"=D{row}-C{row}")
        if steps_sheet.cell(row=row, column=11).value is None:
            steps_sheet.cell(
                row=row,
                column=11,
                value=f'=IF(OR(J{row}="",J{row}=0),"",ABS((E{row}-J{row})/J{row}))',
            )
        for column in range(1, len(headers) + 1):
            cell = steps_sheet.cell(row=row, column=column)
            if cell.has_style is False:
                cell.font = Font(
                    name=FONT_NAME,
                    size=10,
                    color="000000" if column in (5, 11) else "0000FF",
                )
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=column in (6, 8))
        for column in (3, 4, 5, 10):
            steps_sheet.cell(row=row, column=column).number_format = "0.0"
        steps_sheet.cell(row=row, column=11).number_format = "0.0%"

    category_validation = DataValidation(
        type="list",
        formula1='"' + ",".join(CATEGORIES) + '"',
        allow_blank=False,
    )
    steps_sheet.add_data_validation(category_validation)
    category_validation.add(f"G2:G{formula_end_row}")
    for column, width in zip(
        "ABCDEFGHIJK",
        [18, 6, 11, 11, 11, 46, 18, 38, 20, 15, 20],
    ):
        steps_sheet.column_dimensions[column].width = width
    steps_sheet.freeze_panes = "A2"
    steps_sheet.auto_filter.ref = f"A1:K{formula_end_row}"

    summary_sheet = workbook.create_sheet("Synthèse")
    summary_sheet["A1"] = f"Synthèse - {project['nom']} (temps en secondes)"
    summary_sheet["A1"].font = Font(name=FONT_NAME, bold=True, size=14)
    summary_sheet["A2"] = (
        "Temps ESTIMÉS à valider par chronométrage réel. En mode mouvement, seule la présence "
        "de mouvement est mesurée, pas la nature du travail."
    )
    summary_sheet["A2"].font = Font(name=FONT_NAME, italic=True, color="C00000", size=10)
    summary_sheet.merge_cells("A2:J2")
    summary_sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    summary_sheet.row_dimensions[2].height = 32

    summary_headers = [
        "Poste",
        *CATEGORIES,
        "Temps de cycle",
        "Takt",
        "% VA",
        "Écart cycle - takt",
        "Erreur moyenne vs chrono (%)",
    ]
    summary_header_row = 5
    for column, header in enumerate(summary_headers, start=1):
        summary_sheet.cell(row=summary_header_row, column=column, value=header)
    _style_header(summary_sheet, summary_header_row, len(summary_headers))

    first_summary_row = summary_header_row + 1
    for station_index, station in enumerate(stations):
        summary_row = first_summary_row + station_index
        summary_sheet.cell(row=summary_row, column=1, value=_safe_text(station["nom"]))
        for category_index, _category in enumerate(CATEGORIES):
            summary_column = 2 + category_index
            summary_letter = summary_sheet.cell(row=summary_header_row, column=summary_column).column_letter
            summary_sheet.cell(
                row=summary_row,
                column=summary_column,
                value=(
                    f"=SUMIFS('Étapes'!$E$2:$E${formula_end_row},"
                    f"'Étapes'!$A$2:$A${formula_end_row},$A{summary_row},"
                    f"'Étapes'!$G$2:$G${formula_end_row},{summary_letter}${summary_header_row})"
                ),
            )
        summary_sheet.cell(row=summary_row, column=6, value=f"=SUM(B{summary_row}:E{summary_row})")
        summary_sheet.cell(row=summary_row, column=7, value=float(project["takt_s"]))
        summary_sheet.cell(row=summary_row, column=8, value=f"=IF(F{summary_row}=0,0,B{summary_row}/F{summary_row})")
        summary_sheet.cell(row=summary_row, column=9, value=f"=IF(G{summary_row}=0,0,F{summary_row}-G{summary_row})")
        summary_sheet.cell(
            row=summary_row,
            column=10,
            value=f'=IFERROR(AVERAGEIF(\'Étapes\'!$A$2:$A${formula_end_row},$A{summary_row},\'Étapes\'!$K$2:$K${formula_end_row}),"")',
        )
        for column in range(1, 11):
            cell = summary_sheet.cell(row=summary_row, column=column)
            cell.font = Font(name=FONT_NAME, size=10, color="0000FF" if column == 7 else "000000")
            cell.number_format = "0.0%" if column in (8, 10) else "0.0"

    end_summary_row = max(first_summary_row + len(stations) - 1, first_summary_row)
    summary_sheet.column_dimensions["A"].width = 22
    for column in "BCDEFGHIJ":
        summary_sheet.column_dimensions[column].width = 19
    summary_sheet.row_dimensions[summary_header_row].height = 34

    if stations:
        bar_chart = BarChart()
        bar_chart.type = "col"
        bar_chart.grouping = "stacked"
        bar_chart.overlap = 100
        bar_chart.title = "Yamazumi - temps par poste"
        bar_chart.y_axis.title = "Secondes"
        bar_chart.x_axis.title = "Poste"
        bar_chart.add_data(
            Reference(summary_sheet, min_col=2, max_col=5, min_row=summary_header_row, max_row=end_summary_row),
            titles_from_data=True,
        )
        bar_chart.set_categories(
            Reference(summary_sheet, min_col=1, min_row=first_summary_row, max_row=end_summary_row)
        )
        for series, category in zip(bar_chart.series, CATEGORIES):
            series.graphicalProperties = GraphicalProperties(solidFill=CATEGORY_COLORS[category])
        bar_chart.height = 10
        bar_chart.width = 20

        if float(project["takt_s"]) > 0:
            takt_chart = LineChart()
            takt_chart.add_data(
                Reference(summary_sheet, min_col=7, min_row=summary_header_row, max_row=end_summary_row),
                titles_from_data=True,
            )
            takt_chart.set_categories(
                Reference(summary_sheet, min_col=1, min_row=first_summary_row, max_row=end_summary_row)
            )
            takt_chart.series[0].graphicalProperties.line.solidFill = "000000"
            takt_chart.series[0].graphicalProperties.line.dashStyle = "dash"
            takt_chart.series[0].marker.symbol = "none"
            bar_chart += takt_chart
        summary_sheet.add_chart(bar_chart, f"A{end_summary_row + 3}")

    notes_sheet = workbook.create_sheet("Notes")
    notes = [
        "Kaizen Studio - étude de temps assistée par vidéo",
        f"Classeur généré le {datetime.now():%d/%m/%Y à %H:%M}",
        "",
        "Les temps issus de l'analyse vidéo sont des ESTIMATIONS à valider par chronométrage réel.",
        "En mode motion, l'application mesure uniquement la présence de mouvement et ne reconnaît pas les gestes.",
        "Les durées sont calculées par formule Excel (fin - début). La synthèse utilise SUMIFS par catégorie.",
        "Les plages de formules comprennent une marge de 60 lignes pour permettre l'ajout d'étapes.",
        "Catégories : VA = valeur ajoutée ; NVA nécessaire = indispensable mais sans valeur client directe ; "
        "Gaspillage = activité à éliminer ; Non classé = à analyser.",
        "RGPD : informer les opérateurs et valider le cadre social et de protection des données avant de filmer.",
    ]
    for row, note in enumerate(notes, start=1):
        cell = notes_sheet.cell(row=row, column=1, value=note)
        cell.font = Font(name=FONT_NAME, bold=row == 1, size=12 if row == 1 else 10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    notes_sheet.column_dimensions["A"].width = 110
    cycle_table_title_row = len(notes) + 2
    cycle_headers = ["Poste", "N° cycle", "Début (s)", "Fin (s)", "Durée (s)", "Origine"]
    for column, header in enumerate(cycle_headers, start=1):
        notes_sheet.cell(row=cycle_table_title_row, column=column, value=header)
    _style_header(notes_sheet, cycle_table_title_row, len(cycle_headers))
    cycle_row = cycle_table_title_row + 1
    for station in stations:
        for cycle_index, cycle in enumerate(station.get("cycles", []), start=1):
            notes_sheet.append(
                [
                    _safe_text(station["nom"]),
                    cycle_index,
                    cycle["debut_s"],
                    cycle["fin_s"],
                    f"=D{cycle_row}-C{cycle_row}",
                    "Proposition automatique vérifiée" if cycle.get("source") == "auto" else "Délimitation manuelle",
                ]
            )
            for column in range(1, len(cycle_headers) + 1):
                cell = notes_sheet.cell(row=cycle_row, column=column)
                cell.font = Font(
                    name=FONT_NAME,
                    size=10,
                    color="000000" if column == 5 else "0000FF",
                )
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            for column in (3, 4, 5):
                notes_sheet.cell(row=cycle_row, column=column).number_format = "0.0"
            cycle_row += 1
    for column, width in zip("BCDEF", [12, 14, 14, 14, 34]):
        notes_sheet.column_dimensions[column].width = width

    workbook.move_sheet(summary_sheet, offset=-1)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def export_project_workbook(
    project: dict[str, Any],
    stations: list[dict[str, Any]],
    target_path: str | Path,
) -> tuple[Path, str | None]:
    """Écrit l'Excel ; en cas de verrouillage, réessaie vers un nom horodaté."""
    target_path = Path(target_path)
    try:
        return build_project_workbook(project, stations, target_path), None
    except OSError as first_error:
        timestamped_path = target_path.with_name(
            f"{target_path.stem}_{datetime.now():%Y%m%d_%H%M%S}{target_path.suffix}"
        )
        try:
            return build_project_workbook(project, stations, timestamped_path), (
                f"Écriture impossible sur {target_path.name} ({first_error}). "
                f"Le classeur a été enregistré sous {timestamped_path.name}."
            )
        except OSError as retry_error:
            raise OSError(
                f"Échec de l'export Excel vers {target_path} puis {timestamped_path}: {retry_error}"
            ) from retry_error
