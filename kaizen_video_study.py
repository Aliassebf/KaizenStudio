#!/usr/bin/env python
"""
kaizen_video_study.py - Prototype d'étude de temps assistée par IA (vidéo -> Excel + Yamazumi).

Pipeline : vidéo -> images horodatées (OpenCV) -> modèle de vision Azure OpenAI
           -> étapes (début/fin/catégorie) -> classeur Excel (formules + graphique Yamazumi).

ATTENTION : les temps sont des ESTIMATIONS issues d'images échantillonnées.
Ils doivent être validés par un chronométrage réel (colonne "Chrono réel" du classeur).

Exemples :
    python kaizen_video_study.py poste1.mp4 --takt 60 --mock            # test sans API
    python kaizen_video_study.py poste1.mp4 poste2.mp4 --takt 60         # analyse réelle
    python kaizen_video_study.py poste1.mp4 --start 12 --end 75 --fps 2  # un seul cycle
"""
import argparse
import base64
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import cv2
import requests
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

CATEGORIES = ["VA", "NVA nécessaire", "Gaspillage", "Non classé"]
CAT_COLORS = {"VA": "2E7D32", "NVA nécessaire": "F9A825", "Gaspillage": "C62828", "Non classé": "9E9E9E"}
FONT = "Arial"

SYSTEM_PROMPT = """Tu es un ingénieur industrialisation expert Lean / Toyota Production System \
et en études de temps et mouvements. Tu analyses des images extraites d'une vidéo d'un poste de travail.
Chaque image porte un horodatage incrusté en haut à gauche (t=XX.Xs).

Mission : découper l'activité observée en opérations élémentaires successives et, pour chacune, donner :
- start_s / end_s : début et fin en secondes, en t'appuyant UNIQUEMENT sur les horodatages visibles ;
- description : verbe d'action + objet (ex. « Saisir le circuit imprimé dans le bac ») ;
- category : "VA" (transforme le produit, valeur pour le client), "NVA nécessaire" \
(déplacement, saisie, contrôle, manutention indispensables en l'état du process) ou "Gaspillage" \
(attente, recherche, mouvement inutile, reprise, surstock) ;
- reason : justification courte de la catégorie ;
- confidence : "haute", "moyenne" ou "faible".

Règles : ne rien inventer, ne pas deviner ce qui n'est pas visible ; en cas de doute, confidence "faible". \
Les étapes doivent être consécutives et sans chevauchement. Si une étape continue après la dernière \
image du lot, termine-la à l'horodatage de la dernière image.
Réponds UNIQUEMENT avec un objet JSON : {"steps": [{"start_s": 0.0, "end_s": 3.0, "description": "...", \
"category": "VA", "reason": "...", "confidence": "moyenne"}]}"""


# --------------------------------------------------------------------------- vidéo
def extract_frames(video_path, fps, width, t_start=0.0, t_end=None):
    """Retourne (liste de (t, jpeg_bytes), durée_fenêtre_s)."""
    cap = cv2.VideoCapture(str(video_path))
    if not cap.isOpened():
        sys.exit(f"Impossible d'ouvrir la vidéo : {video_path}")
    src_fps = cap.get(cv2.CAP_PROP_FPS) or 25.0
    total = cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0
    duration = total / src_fps if total else 0
    t_end = min(t_end, duration) if (t_end and duration) else (t_end or duration)
    if t_end <= t_start:
        sys.exit(f"Fenêtre invalide pour {video_path} (début {t_start}s, fin {t_end}s).")

    frames, t = [], t_start
    step = 1.0 / fps
    while t < t_end:
        cap.set(cv2.CAP_PROP_POS_MSEC, t * 1000)
        ok, img = cap.read()
        if not ok:
            break
        h, w = img.shape[:2]
        if w > width:
            img = cv2.resize(img, (width, int(h * width / w)))
        label = f"t={t:.1f}s"
        cv2.rectangle(img, (0, 0), (150, 30), (0, 0, 0), -1)
        cv2.putText(img, label, (6, 22), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (255, 255, 255), 2)
        ok, buf = cv2.imencode(".jpg", img, [cv2.IMWRITE_JPEG_QUALITY, 80])
        if ok:
            frames.append((round(t, 2), buf.tobytes()))
        t += step
    cap.release()
    return frames, t_end - t_start


def chunk_frames(frames, chunk_seconds):
    chunks, current, limit = [], [], None
    for t, img in frames:
        if limit is None:
            limit = t + chunk_seconds
        if t >= limit:
            chunks.append(current)
            current, limit = [], t + chunk_seconds
        current.append((t, img))
    if current:
        chunks.append(current)
    return chunks


# --------------------------------------------------------------------------- modèle
class AzureVision:
    def __init__(self, detail="auto", max_retries=5):
        self.endpoint = os.environ.get("AZURE_OPENAI_ENDPOINT", "").rstrip("/")
        self.deployment = os.environ.get("AZURE_OPENAI_DEPLOYMENT", "")
        self.api_version = os.environ.get("AZURE_OPENAI_API_VERSION", "2024-10-21")
        self.api_key = os.environ.get("AZURE_OPENAI_API_KEY", "")
        self.detail, self.max_retries = detail, max_retries
        if not self.endpoint or not self.deployment:
            sys.exit("Définir AZURE_OPENAI_ENDPOINT et AZURE_OPENAI_DEPLOYMENT (voir .env.sample), ou utiliser --mock.")

    def _headers(self):
        if self.api_key:
            return {"api-key": self.api_key, "Content-Type": "application/json"}
        try:  # authentification Entra ID (az login / identité managée)
            from azure.identity import DefaultAzureCredential
            token = DefaultAzureCredential().get_token("https://cognitiveservices.azure.com/.default").token
            return {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
        except ImportError:
            sys.exit("Pas de AZURE_OPENAI_API_KEY et azure-identity n'est pas installé.")

    def analyze(self, chunk, context_hint):
        content = [{"type": "text", "text": (
            f"Lot d'images de t={chunk[0][0]:.1f}s à t={chunk[-1][0]:.1f}s. {context_hint}")}]
        for t, img in chunk:
            b64 = base64.b64encode(img).decode()
            content.append({"type": "image_url",
                            "image_url": {"url": f"data:image/jpeg;base64,{b64}", "detail": self.detail}})
        body = {"messages": [{"role": "system", "content": SYSTEM_PROMPT},
                             {"role": "user", "content": content}],
                "temperature": 0, "response_format": {"type": "json_object"}, "max_tokens": 2000}
        url = (f"{self.endpoint}/openai/deployments/{self.deployment}/chat/completions"
               f"?api-version={self.api_version}")
        for attempt in range(self.max_retries):
            r = requests.post(url, headers=self._headers(), json=body, timeout=180)
            if r.status_code == 200:
                text = r.json()["choices"][0]["message"]["content"]
                return json.loads(text).get("steps", [])
            if r.status_code in (429, 500, 502, 503, 504):
                wait = int(r.headers.get("Retry-After", 2 ** (attempt + 1)))
                print(f"  API {r.status_code}, nouvelle tentative dans {wait}s...")
                time.sleep(wait)
                continue
            sys.exit(f"Erreur API {r.status_code} : {r.text[:500]}")
        sys.exit("Échec de l'appel API après plusieurs tentatives.")


def mock_steps(chunk):
    """Étapes factices proportionnelles à la durée du lot (test du pipeline uniquement)."""
    t0, t1 = chunk[0][0], chunk[-1][0] + 1.0
    template = [("Saisir la pièce dans le bac", "NVA nécessaire", 0.15), ("Positionner la pièce sur le poste", "VA", 0.20),
                ("Visser / assembler", "VA", 0.30), ("Chercher un outil", "Gaspillage", 0.10),
                ("Contrôler visuellement", "NVA nécessaire", 0.15), ("Déposer la pièce finie", "NVA nécessaire", 0.10)]
    steps, cur = [], t0
    for desc, cat, share in template:
        end = min(t1, cur + share * (t1 - t0))
        steps.append({"start_s": cur, "end_s": end, "description": f"[MOCK] {desc}", "category": cat,
                      "reason": "Donnée factice (mode --mock)", "confidence": "faible"})
        cur = end
    return steps


# --------------------------------------------------------------------------- fusion
def norm_cat(cat):
    c = (cat or "").strip().lower()
    if c.startswith("va") or "valeur ajout" in c and "non" not in c:
        return "VA"
    if "nécess" in c or "necess" in c or c.startswith("nva") or c.startswith("nnva"):
        return "NVA nécessaire"
    if "gaspill" in c or "waste" in c or "muda" in c:
        return "Gaspillage"
    return "Non classé"


def merge_steps(raw, t_start, t_end, gap_tol):
    """Trie, supprime les chevauchements, fusionne les étapes identiques adjacentes, comble les trous."""
    clean = []
    for s in raw:
        try:
            a, b = float(s["start_s"]), float(s["end_s"])
        except (KeyError, TypeError, ValueError):
            continue
        a, b = max(a, t_start), min(b, t_end)
        if b - a <= 0:
            continue
        clean.append({"start": a, "end": b, "desc": str(s.get("description", "")).strip() or "(sans description)",
                      "cat": norm_cat(s.get("category")), "reason": str(s.get("reason", "")).strip(),
                      "conf": str(s.get("confidence", "")).strip()})
    clean.sort(key=lambda x: x["start"])

    out = []
    for s in clean:
        if out:
            prev = out[-1]
            s["start"] = max(s["start"], prev["end"])
            if s["end"] <= s["start"]:
                continue
            same = prev["desc"].lower() == s["desc"].lower() and prev["cat"] == s["cat"]
            if same and s["start"] - prev["end"] <= gap_tol:
                prev["end"] = s["end"]
                continue
            gap = s["start"] - prev["end"]
            if 0 < gap <= gap_tol:
                prev["end"] = s["start"]
            elif gap > gap_tol:
                out.append({"start": prev["end"], "end": s["start"], "desc": "Non décrit (aucune étape renvoyée)",
                            "cat": "Non classé", "reason": "Trou entre deux étapes", "conf": "faible"})
        elif s["start"] - t_start > gap_tol:
            out.append({"start": t_start, "end": s["start"], "desc": "Non décrit (aucune étape renvoyée)",
                        "cat": "Non classé", "reason": "Début non couvert", "conf": "faible"})
        out.append(s)
    if out and t_end - out[-1]["end"] > gap_tol:
        out.append({"start": out[-1]["end"], "end": t_end, "desc": "Non décrit (aucune étape renvoyée)",
                    "cat": "Non classé", "reason": "Fin non couverte", "conf": "faible"})
    return out


def analyze_video(video, args, model):
    print(f"[{video.name}] extraction des images ({args.fps} i/s)...")
    frames, window = extract_frames(video, args.fps, args.width, args.start, args.end)
    if not frames:
        sys.exit(f"Aucune image extraite de {video}.")
    t_start, t_end = frames[0][0], args.start + window
    chunks = chunk_frames(frames, args.chunk_seconds)
    print(f"[{video.name}] {len(frames)} images, {len(chunks)} lot(s), fenêtre {t_start:.1f}s -> {t_end:.1f}s")

    raw, hint = [], "Premier lot de la vidéo."
    for i, chunk in enumerate(chunks, 1):
        print(f"  lot {i}/{len(chunks)}...")
        steps = mock_steps(chunk) if model is None else model.analyze(chunk, hint)
        raw.extend(steps)
        if steps:
            last = steps[-1]
            hint = (f"Contexte : l'étape précédente était « {last.get('description', '')} » "
                    f"(fin à {last.get('end_s')}s). Continue-la si elle se poursuit.")
    return merge_steps(raw, t_start, t_end, gap_tol=1.5 / args.fps), (t_start, t_end)


# --------------------------------------------------------------------------- Excel
def style_header(ws, row, ncols, color="1F3864"):
    thin = Side(style="thin", color="BFBFBF")
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)


def build_workbook(results, takt, out_path, mock):
    wb = Workbook()
    ws = wb.active
    ws.title = "Étapes"
    heads = ["Poste", "N°", "Début (s)", "Fin (s)", "Durée (s)", "Description", "Catégorie",
             "Justification", "Confiance IA", "Chrono réel (s)", "Écart IA vs réel (%)"]
    ws.append(heads)
    style_header(ws, 1, len(heads))
    thin = Side(style="thin", color="D9D9D9")
    r = 2
    for station, (steps, _w) in results.items():
        for n, s in enumerate(steps, 1):
            ws.append([station, n, round(s["start"], 1), round(s["end"], 1), f"=D{r}-C{r}", s["desc"], s["cat"],
                       s["reason"], s["conf"], None, f'=IF(J{r}="","",(E{r}-J{r})/J{r})'])
            for c in range(1, len(heads) + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = Font(name=FONT, size=10, color="0000FF" if c in (3, 4, 10) else "000000")
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=c in (6, 8))
            ws.cell(row=r, column=3).number_format = ws.cell(row=r, column=4).number_format = "0.0"
            ws.cell(row=r, column=5).number_format = ws.cell(row=r, column=10).number_format = "0.0"
            ws.cell(row=r, column=11).number_format = "0.0%"
            r += 1
    last = max(r - 1, 2)
    dv = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{last}")
    for col, w in zip("ABCDEFGHIJK", [16, 5, 10, 10, 10, 48, 16, 38, 12, 13, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{last}"

    # ---- Synthèse + Yamazumi
    sy = wb.create_sheet("Synthèse")
    sy["A1"] = "Synthèse par poste - temps en secondes"
    sy["A1"].font = Font(name=FONT, bold=True, size=14)
    sy["A2"] = ("ESTIMATIONS IA à valider par chronométrage réel (colonne J de l'onglet Étapes). "
                "Catégories modifiables dans l'onglet Étapes : la synthèse se recalcule.")
    sy["A2"].font = Font(name=FONT, italic=True, color="C00000", size=10)
    if mock:
        sy["A3"] = "MODE TEST (--mock) : données factices, aucune analyse réelle de la vidéo."
        sy["A3"].font = Font(name=FONT, bold=True, color="FFFFFF")
        sy["A3"].fill = PatternFill("solid", start_color="C00000")
    sy.merge_cells("A2:I2")
    sy.merge_cells("A3:I3")
    sy["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sy.row_dimensions[2].height = 28
    hdr = ["Poste"] + CATEGORIES + ["Temps de cycle", "Takt", "% VA", "Écart cycle - takt"]
    for i, h in enumerate(hdr, 1):
        sy.cell(row=5, column=i, value=h)
    style_header(sy, 5, len(hdr))
    row = 6
    for station in results:
        sy.cell(row=row, column=1, value=station)
        for i, cat in enumerate(CATEGORIES):
            col = get_column_letter(2 + i)
            sy.cell(row=row, column=2 + i,
                    value=f"=SUMIFS(Étapes!$E$2:$E${last},Étapes!$A$2:$A${last},$A{row},"
                          f"Étapes!$G$2:$G${last},{col}$5)")
        sy.cell(row=row, column=6, value=f"=SUM(B{row}:E{row})")
        sy.cell(row=row, column=7, value=takt if takt else None)
        sy.cell(row=row, column=8, value=f"=IF(F{row}=0,0,B{row}/F{row})")
        sy.cell(row=row, column=9, value=f'=IF(G{row}="","",F{row}-G{row})')
        for c in range(1, 10):
            cell = sy.cell(row=row, column=c)
            cell.font = Font(name=FONT, size=10, color="0000FF" if c == 7 else "000000")
            cell.number_format = "0.0%" if c == 8 else "0.0"
        row += 1
    end_row = row - 1
    sy.column_dimensions["A"].width = 18
    for col in "BCDEFGHI":
        sy.column_dimensions[col].width = 16
    sy.row_dimensions[5].height = 32

    bar = BarChart()
    bar.type, bar.grouping, bar.overlap = "col", "stacked", 100
    bar.title = "Yamazumi - temps par poste (s)"
    bar.y_axis.title, bar.x_axis.title = "Secondes", "Poste"
    bar.add_data(Reference(sy, min_col=2, max_col=5, min_row=5, max_row=end_row), titles_from_data=True)
    bar.set_categories(Reference(sy, min_col=1, min_row=6, max_row=end_row))
    for series, cat in zip(bar.series, CATEGORIES):
        series.graphicalProperties = GraphicalProperties(solidFill=CAT_COLORS[cat])
    bar.x_axis.delete = bar.y_axis.delete = False
    bar.height, bar.width = 10, 18
    if takt:
        line = LineChart()
        line.add_data(Reference(sy, min_col=7, min_row=5, max_row=end_row), titles_from_data=True)
        line.series[0].graphicalProperties.line.solidFill = "000000"
        line.series[0].graphicalProperties.line.dashStyle = "dash"
        line.series[0].marker.symbol = "none"
        bar += line
    sy.add_chart(bar, f"A{end_row + 3}")

    # ---- Notes
    nt = wb.create_sheet("Notes")
    notes = ["Étude de temps assistée par IA - prototype",
             f"Généré le {datetime.now():%d/%m/%Y %H:%M}",
             "",
             "Méthode : images échantillonnées de la vidéo, horodatées, analysées par un modèle de vision.",
             "Limites : précision temporelle limitée par la fréquence d'échantillonnage ; description et catégorisation "
             "à relire par un ingénieur méthodes.",
             "Validation : renseigner le chrono réel (onglet Étapes, colonne J) pour mesurer l'écart.",
             "Catégories : VA = valeur ajoutée ; NVA nécessaire = non VA mais indispensable aujourd'hui ; "
             "Gaspillage = à éliminer ; Non classé = non couvert / incertain.",
             "RGPD : vérifier le cadre (information des opérateurs, instances représentatives) avant de filmer."]
    for i, t in enumerate(notes, 1):
        nt.cell(row=i, column=1, value=t).font = Font(name=FONT, bold=(i == 1), size=12 if i == 1 else 10)
        nt.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    nt.column_dimensions["A"].width = 110
    wb.move_sheet("Synthèse", offset=-1)
    wb.save(out_path)


# --------------------------------------------------------------------------- main
def main():
    p = argparse.ArgumentParser(description="Étude de temps assistée par IA (vidéo -> Excel + Yamazumi)")
    p.add_argument("videos", nargs="+", type=Path, help="Une vidéo par poste")
    p.add_argument("--stations", nargs="+", help="Noms des postes (défaut : nom du fichier)")
    p.add_argument("--takt", type=float, help="Takt time en secondes (ligne sur le Yamazumi)")
    p.add_argument("--fps", type=float, default=1.0, help="Images analysées par seconde (défaut 1)")
    p.add_argument("--chunk-seconds", type=float, default=20.0, help="Durée d'un lot envoyé au modèle")
    p.add_argument("--width", type=int, default=768, help="Largeur max des images envoyées")
    p.add_argument("--start", type=float, default=0.0, help="Début de la fenêtre analysée (s)")
    p.add_argument("--end", type=float, default=None, help="Fin de la fenêtre analysée (s) - isoler un cycle")
    p.add_argument("--detail", choices=["low", "high", "auto"], default="auto")
    p.add_argument("--out", type=Path, default=None, help="Fichier Excel de sortie")
    p.add_argument("--mock", action="store_true", help="Test du pipeline sans appel API (données factices)")
    args = p.parse_args()

    if args.stations and len(args.stations) != len(args.videos):
        sys.exit("--stations doit avoir autant de noms que de vidéos.")
    stations = args.stations or [v.stem for v in args.videos]
    model = None if args.mock else AzureVision(detail=args.detail)

    results = {}
    for name, video in zip(stations, args.videos):
        steps, window = analyze_video(video, args, model)
        results[name] = (steps, window)
        out_json = video.with_suffix(".steps.json")
        out_json.write_text(json.dumps(steps, ensure_ascii=False, indent=2), encoding="utf-8")

    out = args.out or Path(f"etude_temps_{datetime.now():%Y%m%d_%H%M}.xlsx")
    build_workbook(results, args.takt, out, args.mock)
    print(f"\nTerminé : {out}")


if __name__ == "__main__":
    main()
