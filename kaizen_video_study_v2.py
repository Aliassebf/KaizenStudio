#!/usr/bin/env python
"""
kaizen_video_study_v2.py - Prototype d'étude de temps assistée par vidéo (vidéo -> Excel + Yamazumi).

Trois moteurs (--backend) :
  azure  : images horodatées -> modèle de vision Azure OpenAI -> étapes décrites et catégorisées
  motion : détection locale (OpenCV, CPU, sans IA ni réseau) des phases d'activité / d'arrêt ;
           descriptions et catégories à saisir dans l'Excel
  mock   : données factices pour tester l'installation
Sortie : classeur Excel (Étapes, Synthèse avec Yamazumi, Notes).

ATTENTION : les temps sont des ESTIMATIONS issues d'images échantillonnées.
Ils doivent être validés par un chronométrage réel (colonne "Chrono réel" du classeur).

Exemples :
    python kaizen_video_study_v2.py poste1.mp4 --takt 60 --mock             # test sans API
    python kaizen_video_study_v2.py poste1.mp4 --backend motion --takt 60   # 100 % local, sans IA
    python kaizen_video_study_v2.py poste1.mp4 --start 12 --end 75 --backend motion --roi 0.2,0.1,0.6,0.8
    python kaizen_video_study_v2.py poste1.mp4 --start 12 --end 75 --fps 2   # moteur azure
"""
import argparse
import base64
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
import requests
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

CATEGORIES = ["VA", "NVA nécessaire", "Gaspillage", "Non classé"]
CAT_COLORS = {"VA": "2E7D32", "NVA nécessaire": "F9A825", "Gaspillage": "C62828", "Non classé": "9E9E9E"}
FONT = "Arial"

SYSTEM_PROMPT = """Tu es un ingénieur industrialisation expert Lean / Toyota Production System \
et en études de temps et mouvements. Tu analyses des images extraites d'une vidéo d'un poste de travail.
Chaque image porte un horodatage incrusté en haut à gauche (t=XX.Xs).

Mission : découper l'activité observée en opérations élémentaires successives et, pour chacune, donner :
- start_s / end_s : début et fin en secondes, en t'appuyant UNIQUEMENT sur les horodatages visibles ;
- description : verbe d'action + objet (ex. « Saisir le circuit imprimé dans le bac ») ;
- category : "VA" (transforme le produit, valeur pour le client), "NVA nécessaire" \
(déplacement, saisie, contrôle, manutention indispensables en l'état du process) ou "Gaspillage" \
(attente, recherche, mouvement inutile, reprise, surstock) ;
- reason : justification courte de la catégorie ;
- confidence : "haute", "moyenne" ou "faible".

Règles : ne rien inventer, ne pas deviner ce qui n'est pas visible ; en cas de doute, confidence "faible". \
Les étapes doivent être consécutives et sans chevauchement. Si une étape continue après la dernière \
image du lot, termine-la à l'horodatage de la dernière image.
Réponds UNIQUEMENT avec un objet JSON : {"steps": [{"start_s": 0.0, "end_s": 3.0, "description": "...", \
"category": "VA", "reason": "...", "confidence": "moyenne"}]}"""


# --------------------------------------------------------------------------- vidéo
def extract_frames(video_path, fps, width, t_start=0.0, t_end=None):
    """Retourne (liste de (t, jpeg_bytes), durée_fenêtre_s)."""
    cap = cv2.VideoCapture(str(video_path))
    if not cap.isOpened():
        sys.exit(f"Impossible d'ouvrir la vidéo : {video_path}")
    src_fps = cap.get(cv2.CAP_PROP_FPS) or 25.0
    total = cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0
    duration = total / src_fps if total else 0
    t_end = min(t_end, duration) if (t_end and duration) else (t_end or duration)
    if t_end <= t_start:
        sys.exit(f"Fenêtre invalide pour {video_path} (début {t_start}s, fin {t_end}s).")

    frames, t = [], t_start
    step = 1.0 / fps
    while t < t_end:
        cap.set(cv2.CAP_PROP_POS_MSEC, t * 1000)
        ok, img = cap.read()
        if not ok:
            break
        h, w = img.shape[:2]
        if w > width:
            img = cv2.resize(img, (width, int(h * width / w)))
        label = f"t={t:.1f}s"
        cv2.rectangle(img, (0, 0), (150, 30), (0, 0, 0), -1)
        cv2.putText(img, label, (6, 22), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (255, 255, 255), 2)
        ok, buf = cv2.imencode(".jpg", img, [cv2.IMWRITE_JPEG_QUALITY, 80])
        if ok:
            frames.append((round(t, 2), buf.tobytes()))
        t += step
    cap.release()
    return frames, t_end - t_start


def chunk_frames(frames, chunk_seconds):
    chunks, current, limit = [], [], None
    for t, img in frames:
        if limit is None:
            limit = t + chunk_seconds
        if t >= limit:
            chunks.append(current)
            current, limit = [], t + chunk_seconds
        current.append((t, img))
    if current:
        chunks.append(current)
    return chunks


# --------------------------------------------------------------------------- modèle
class AzureVision:
    def __init__(self, detail="auto", max_retries=5):
        self.endpoint = os.environ.get("AZURE_OPENAI_ENDPOINT", "").rstrip("/")
        self.deployment = os.environ.get("AZURE_OPENAI_DEPLOYMENT", "")
        self.api_version = os.environ.get("AZURE_OPENAI_API_VERSION", "2024-10-21")
        self.api_key = os.environ.get("AZURE_OPENAI_API_KEY", "")
        self.detail, self.max_retries = detail, max_retries
        if not self.endpoint or not self.deployment:
            sys.exit("Définir AZURE_OPENAI_ENDPOINT et AZURE_OPENAI_DEPLOYMENT (voir .env.sample), "
                     "ou utiliser --backend motion (local, sans IA) / --mock.")

    def _headers(self):
        if self.api_key:
            return {"api-key": self.api_key, "Content-Type": "application/json"}
        try:  # authentification Entra ID (az login / identité managée)
            from azure.identity import DefaultAzureCredential
            token = DefaultAzureCredential().get_token("https://cognitiveservices.azure.com/.default").token
            return {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
        except ImportError:
            sys.exit("Pas de AZURE_OPENAI_API_KEY et azure-identity n'est pas installé.")

    def analyze(self, chunk, context_hint):
        content = [{"type": "text", "text": (
            f"Lot d'images de t={chunk[0][0]:.1f}s à t={chunk[-1][0]:.1f}s. {context_hint}")}]
        for t, img in chunk:
            b64 = base64.b64encode(img).decode()
            content.append({"type": "image_url",
                            "image_url": {"url": f"data:image/jpeg;base64,{b64}", "detail": self.detail}})
        body = {"messages": [{"role": "system", "content": SYSTEM_PROMPT},
                             {"role": "user", "content": content}],
                "temperature": 0, "response_format": {"type": "json_object"}, "max_tokens": 2000}
        url = (f"{self.endpoint}/openai/deployments/{self.deployment}/chat/completions"
               f"?api-version={self.api_version}")
        for attempt in range(self.max_retries):
            r = requests.post(url, headers=self._headers(), json=body, timeout=180)
            if r.status_code == 200:
                text = r.json()["choices"][0]["message"]["content"]
                return json.loads(text).get("steps", [])
            if r.status_code in (429, 500, 502, 503, 504):
                wait = int(r.headers.get("Retry-After", 2 ** (attempt + 1)))
                print(f"  API {r.status_code}, nouvelle tentative dans {wait}s...")
                time.sleep(wait)
                continue
            sys.exit(f"Erreur API {r.status_code} : {r.text[:500]}")
        sys.exit("Échec de l'appel API après plusieurs tentatives.")


def mock_steps(chunk):
    """Étapes factices proportionnelles à la durée du lot (test du pipeline uniquement)."""
    t0, t1 = chunk[0][0], chunk[-1][0] + 1.0
    template = [("Saisir la pièce dans le bac", "NVA nécessaire", 0.15), ("Positionner la pièce sur le poste", "VA", 0.20),
                ("Visser / assembler", "VA", 0.30), ("Chercher un outil", "Gaspillage", 0.10),
                ("Contrôler visuellement", "NVA nécessaire", 0.15), ("Déposer la pièce finie", "NVA nécessaire", 0.10)]
    steps, cur = [], t0
    for desc, cat, share in template:
        end = min(t1, cur + share * (t1 - t0))
        steps.append({"start_s": cur, "end_s": end, "description": f"[MOCK] {desc}", "category": cat,
                      "reason": "Donnée factice (mode --mock)", "confidence": "faible"})
        cur = end
    return steps


# --------------------------------------------------------------------------- fusion
def norm_cat(cat):
    c = (cat or "").strip().lower()
    if c.startswith("va") or "valeur ajout" in c and "non" not in c:
        return "VA"
    if "nécess" in c or "necess" in c or c.startswith("nva") or c.startswith("nnva"):
        return "NVA nécessaire"
    if "gaspill" in c or "waste" in c or "muda" in c:
        return "Gaspillage"
    return "Non classé"


def merge_steps(raw, t_start, t_end, gap_tol):
    """Trie, supprime les chevauchements, fusionne les étapes identiques adjacentes, comble les trous."""
    clean = []
    for s in raw:
        try:
            a, b = float(s["start_s"]), float(s["end_s"])
        except (KeyError, TypeError, ValueError):
            continue
        a, b = max(a, t_start), min(b, t_end)
        if b - a <= 0:
            continue
        clean.append({"start": a, "end": b, "desc": str(s.get("description", "")).strip() or "(sans description)",
                      "cat": norm_cat(s.get("category")), "reason": str(s.get("reason", "")).strip(),
                      "conf": str(s.get("confidence", "")).strip()})
    clean.sort(key=lambda x: x["start"])

    out = []
    for s in clean:
        if out:
            prev = out[-1]
            s["start"] = max(s["start"], prev["end"])
            if s["end"] <= s["start"]:
                continue
            same = prev["desc"].lower() == s["desc"].lower() and prev["cat"] == s["cat"]
            if same and s["start"] - prev["end"] <= gap_tol:
                prev["end"] = s["end"]
                continue
            gap = s["start"] - prev["end"]
            if 0 < gap <= gap_tol:
                prev["end"] = s["start"]
            elif gap > gap_tol:
                out.append({"start": prev["end"], "end": s["start"], "desc": "Non décrit (aucune étape renvoyée)",
                            "cat": "Non classé", "reason": "Trou entre deux étapes", "conf": "faible"})
        elif s["start"] - t_start > gap_tol:
            out.append({"start": t_start, "end": s["start"], "desc": "Non décrit (aucune étape renvoyée)",
                        "cat": "Non classé", "reason": "Début non couvert", "conf": "faible"})
        out.append(s)
    if out and t_end - out[-1]["end"] > gap_tol:
        out.append({"start": out[-1]["end"], "end": t_end, "desc": "Non décrit (aucune étape renvoyée)",
                    "cat": "Non classé", "reason": "Fin non couverte", "conf": "faible"})
    return out


# --------------------------------------------------------------------------- mouvement (local, sans IA)
def motion_scores(video, fps, t_start, t_end, roi, pix_thr, blur):
    """Score de mouvement = part des pixels dont la luminosité change entre deux images échantillonnées."""
    cap = cv2.VideoCapture(str(video))
    if not cap.isOpened():
        sys.exit(f"Impossible d'ouvrir la vidéo : {video}")
    src_fps = cap.get(cv2.CAP_PROP_FPS) or 25.0
    total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
    duration = total / src_fps if total else 0
    t_end = min(t_end, duration) if (t_end and duration) else (t_end or duration)
    if t_end <= t_start:
        sys.exit(f"Fenêtre invalide pour {video} (début {t_start}s, fin {t_end}s).")
    step = max(1, round(src_fps / fps))
    first, last = int(t_start * src_fps), int(t_end * src_fps)
    cap.set(cv2.CAP_PROP_POS_FRAMES, first)
    times, scores, prev, i = [], [], None, first
    while i < last:
        if not cap.grab():
            break
        if (i - first) % step == 0:
            ok, img = cap.retrieve()
            if ok:
                h, w = img.shape[:2]
                img = cv2.resize(img, (320, max(1, int(h * 320 / w))), interpolation=cv2.INTER_AREA)
                gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
                if roi:
                    gh, gw = gray.shape
                    x, y, rw, rh = roi
                    gray = gray[int(y * gh):int((y + rh) * gh), int(x * gw):int((x + rw) * gw)]
                gray = cv2.GaussianBlur(gray, (blur, blur), 0)
                if prev is not None:
                    scores.append(float((cv2.absdiff(gray, prev) > pix_thr).mean()))
                    times.append(i / src_fps)
                prev = gray
        i += 1
    cap.release()
    if not scores:
        sys.exit("Aucune image exploitable pour la détection de mouvement.")
    return np.array(times), np.array(scores), t_start, t_end, step / src_fps


def parse_roi(text):
    if not text:
        return None
    try:
        x, y, w, h = [float(v) for v in text.split(",")]
    except ValueError:
        sys.exit("--roi attend 4 fractions entre 0 et 1 : x,y,largeur,hauteur (ex. 0.2,0.1,0.6,0.8)")
    if not (0 <= x < 1 and 0 <= y < 1 and 0 < w <= 1 and 0 < h <= 1 and x + w <= 1.001 and y + h <= 1.001):
        sys.exit("--roi hors de l'image : x+largeur et y+hauteur doivent rester <= 1")
    return x, y, w, h


def hysteresis(scores, thr_on, thr_off):
    state, labels = bool(scores[0] >= thr_on), []
    for sc in scores:
        if not state and sc >= thr_on:
            state = True
        elif state and sc < thr_off:
            state = False
        labels.append(state)
    return labels


def merge_same_label(phases):
    out = []
    for p in phases:
        if out and out[-1]["active"] == p["active"]:
            out[-1]["end"] = p["end"]
        else:
            out.append(dict(p))
    return out


def absorb_short(phases, min_len):
    phases = merge_same_label(phases)
    while len(phases) > 1:
        short = [k for k, p in enumerate(phases) if p["end"] - p["start"] < min_len]
        if not short:
            break
        k = min(short, key=lambda j: phases[j]["end"] - phases[j]["start"])
        left = phases[k - 1] if k > 0 else None
        right = phases[k + 1] if k < len(phases) - 1 else None
        if right is None or (left is not None and left["end"] - left["start"] >= right["end"] - right["start"]):
            left["end"] = phases[k]["end"]
        else:
            right["start"] = phases[k]["start"]
        del phases[k]
        phases = merge_same_label(phases)
    return phases


def save_motion_plot(path, times, smooth, thr_on, phases, t_start, t_end):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return False
    fig, ax = plt.subplots(figsize=(11, 3.6))
    for p in phases:
        ax.axvspan(p["start"], p["end"], color="#2E7D32" if p["active"] else "#BDBDBD", alpha=0.25)
    ax.plot(times, smooth * 100, color="#1F3864", lw=1.2)
    ax.axhline(thr_on * 100, color="#C62828", ls="--", lw=1, label=f"seuil d'activité ({thr_on * 100:.2f} %)")
    ax.set_xlim(t_start, t_end)
    ax.set_xlabel("Temps (s)")
    ax.set_ylabel("Pixels en mouvement (%)")
    ax.set_title("Mouvement détecté (vert = activité, gris = arrêt) - à comparer à la vidéo")
    ax.legend(loc="upper right")
    fig.tight_layout()
    fig.savefig(path, dpi=110)
    plt.close(fig)
    return True


def analyze_motion_video(video, args):
    roi = parse_roi(args.roi)
    print(f"[{video.name}] détection de mouvement ({args.motion_fps} i/s, CPU)...")
    times, scores, t_start, t_end, dt = motion_scores(
        video, args.motion_fps, args.start, args.end, roi, args.pixel_threshold, args.blur)
    k = max(1, int(round(args.smooth / dt)))
    smooth = np.convolve(scores, np.ones(k) / k, mode="same") if k > 1 else scores.copy()
    if args.motion_threshold is not None:
        thr_on = args.motion_threshold / 100.0
    else:
        thr_on = max(0.002, 0.2 * float(np.percentile(smooth, 90)))
    thr_off = 0.6 * thr_on
    labels = hysteresis(smooth, thr_on, thr_off)

    edges = list(times - dt) + [times[-1]]
    edges[0] = t_start
    phases = [{"start": edges[i], "end": edges[i + 1], "active": lab} for i, lab in enumerate(labels)]
    phases[-1]["end"] = t_end
    phases = absorb_short(phases, args.min_phase)

    steps = []
    for p in phases:
        sel = (times > p["start"]) & (times <= p["end"] + 1e-9)
        mean = float(smooth[sel].mean() * 100) if sel.any() else 0.0
        steps.append({
            "start": p["start"], "end": p["end"],
            "desc": "Activité détectée - à décrire" if p["active"] else "Arrêt (aucun mouvement détecté) - à décrire",
            "cat": "Non classé",
            "reason": f"Mouvement moyen {mean:.2f} % des pixels (seuil {thr_on * 100:.2f} %)",
            "conf": "mesurée"})
    png = video.with_suffix(".motion.png")
    if save_motion_plot(png, times, smooth, thr_on, phases, t_start, t_end):
        print(f"[{video.name}] courbe de mouvement : {png}")
    n_act = sum(1 for p in phases if p["active"])
    print(f"[{video.name}] {len(phases)} phases ({n_act} d'activité), seuil {thr_on * 100:.2f} %, "
          f"fenêtre {t_start:.1f}s -> {t_end:.1f}s")
    return steps, (t_start, t_end)


# --------------------------------------------------------------------------- moteur vision (Azure / mock)
def analyze_video(video, args, model):
    print(f"[{video.name}] extraction des images ({args.fps} i/s)...")
    frames, window = extract_frames(video, args.fps, args.width, args.start, args.end)
    if not frames:
        sys.exit(f"Aucune image extraite de {video}.")
    t_start, t_end = frames[0][0], args.start + window
    chunks = chunk_frames(frames, args.chunk_seconds)
    print(f"[{video.name}] {len(frames)} images, {len(chunks)} lot(s), fenêtre {t_start:.1f}s -> {t_end:.1f}s")

    raw, hint = [], "Premier lot de la vidéo."
    for i, chunk in enumerate(chunks, 1):
        print(f"  lot {i}/{len(chunks)}...")
        steps = mock_steps(chunk) if model is None else model.analyze(chunk, hint)
        raw.extend(steps)
        if steps:
            last = steps[-1]
            hint = (f"Contexte : l'étape précédente était « {last.get('description', '')} » "
                    f"(fin à {last.get('end_s')}s). Continue-la si elle se poursuit.")
    return merge_steps(raw, t_start, t_end, gap_tol=1.5 / args.fps), (t_start, t_end)


# --------------------------------------------------------------------------- Excel
def style_header(ws, row, ncols, color="1F3864"):
    thin = Side(style="thin", color="BFBFBF")
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)


def build_workbook(results, takt, out_path, banner=None, method_note=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Étapes"
    heads = ["Poste", "N°", "Début (s)", "Fin (s)", "Durée (s)", "Description", "Catégorie",
             "Justification", "Confiance", "Chrono réel (s)", "Écart auto vs réel (%)"]
    ws.append(heads)
    style_header(ws, 1, len(heads))
    thin = Side(style="thin", color="D9D9D9")
    r = 2
    for station, (steps, _w) in results.items():
        for n, s in enumerate(steps, 1):
            ws.append([station, n, round(s["start"], 1), round(s["end"], 1), f"=D{r}-C{r}", s["desc"], s["cat"],
                       s["reason"], s["conf"], None, f'=IF(J{r}="","",(E{r}-J{r})/J{r})'])
            for c in range(1, len(heads) + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = Font(name=FONT, size=10, color="0000FF" if c in (3, 4, 10) else "000000")
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=c in (6, 8))
            ws.cell(row=r, column=3).number_format = ws.cell(row=r, column=4).number_format = "0.0"
            ws.cell(row=r, column=5).number_format = ws.cell(row=r, column=10).number_format = "0.0"
            ws.cell(row=r, column=11).number_format = "0.0%"
            r += 1
    last = max(r - 1, 2)
    dv = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{last}")
    for col, w in zip("ABCDEFGHIJK", [16, 5, 10, 10, 10, 48, 16, 38, 12, 13, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{last}"

    # ---- Synthèse + Yamazumi
    sy = wb.create_sheet("Synthèse")
    sy["A1"] = "Synthèse par poste - temps en secondes"
    sy["A1"].font = Font(name=FONT, bold=True, size=14)
    sy["A2"] = ("Temps ESTIMÉS automatiquement, à valider par chronométrage réel (colonne J de l'onglet Étapes). "
                "Catégories modifiables dans l'onglet Étapes : la synthèse se recalcule.")
    sy["A2"].font = Font(name=FONT, italic=True, color="C00000", size=10)
    if banner:
        sy["A3"] = banner[0]
        sy["A3"].font = Font(name=FONT, bold=True, color="FFFFFF")
        sy["A3"].fill = PatternFill("solid", start_color=banner[1])
        sy["A3"].alignment = Alignment(wrap_text=True, vertical="center")
        sy.row_dimensions[3].height = 32
    sy.merge_cells("A2:I2")
    sy.merge_cells("A3:I3")
    sy["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sy.row_dimensions[2].height = 28
    hdr = ["Poste"] + CATEGORIES + ["Temps de cycle", "Takt", "% VA", "Écart cycle - takt"]
    for i, h in enumerate(hdr, 1):
        sy.cell(row=5, column=i, value=h)
    style_header(sy, 5, len(hdr))
    row = 6
    for station in results:
        sy.cell(row=row, column=1, value=station)
        for i, cat in enumerate(CATEGORIES):
            col = get_column_letter(2 + i)
            sy.cell(row=row, column=2 + i,
                    value=f"=SUMIFS(Étapes!$E$2:$E${last},Étapes!$A$2:$A${last},$A{row},"
                          f"Étapes!$G$2:$G${last},{col}$5)")
        sy.cell(row=row, column=6, value=f"=SUM(B{row}:E{row})")
        sy.cell(row=row, column=7, value=takt if takt else None)
        sy.cell(row=row, column=8, value=f"=IF(F{row}=0,0,B{row}/F{row})")
        sy.cell(row=row, column=9, value=f'=IF(G{row}="","",F{row}-G{row})')
        for c in range(1, 10):
            cell = sy.cell(row=row, column=c)
            cell.font = Font(name=FONT, size=10, color="0000FF" if c == 7 else "000000")
            cell.number_format = "0.0%" if c == 8 else "0.0"
        row += 1
    end_row = row - 1
    sy.column_dimensions["A"].width = 18
    for col in "BCDEFGHI":
        sy.column_dimensions[col].width = 16
    sy.row_dimensions[5].height = 32

    bar = BarChart()
    bar.type, bar.grouping, bar.overlap = "col", "stacked", 100
    bar.title = "Yamazumi - temps par poste (s)"
    bar.y_axis.title, bar.x_axis.title = "Secondes", "Poste"
    bar.add_data(Reference(sy, min_col=2, max_col=5, min_row=5, max_row=end_row), titles_from_data=True)
    bar.set_categories(Reference(sy, min_col=1, min_row=6, max_row=end_row))
    for series, cat in zip(bar.series, CATEGORIES):
        series.graphicalProperties = GraphicalProperties(solidFill=CAT_COLORS[cat])
    bar.x_axis.delete = bar.y_axis.delete = False
    bar.height, bar.width = 10, 18
    if takt:
        line = LineChart()
        line.add_data(Reference(sy, min_col=7, min_row=5, max_row=end_row), titles_from_data=True)
        line.series[0].graphicalProperties.line.solidFill = "000000"
        line.series[0].graphicalProperties.line.dashStyle = "dash"
        line.series[0].marker.symbol = "none"
        bar += line
    sy.add_chart(bar, f"A{end_row + 3}")

    # ---- Notes
    nt = wb.create_sheet("Notes")
    notes = ["Étude de temps assistée par vidéo - prototype",
             f"Généré le {datetime.now():%d/%m/%Y %H:%M}",
             "",
             method_note or "Méthode : images échantillonnées de la vidéo, horodatées, analysées par un modèle de vision.",
             "Limites : précision temporelle limitée par la fréquence d'échantillonnage ; description et catégorisation "
             "à relire par un ingénieur méthodes.",
             "Validation : renseigner le chrono réel (onglet Étapes, colonne J) pour mesurer l'écart.",
             "Catégories : VA = valeur ajoutée ; NVA nécessaire = non VA mais indispensable aujourd'hui ; "
             "Gaspillage = à éliminer ; Non classé = non couvert / incertain.",
             "RGPD : vérifier le cadre (information des opérateurs, instances représentatives) avant de filmer."]
    for i, t in enumerate(notes, 1):
        nt.cell(row=i, column=1, value=t).font = Font(name=FONT, bold=(i == 1), size=12 if i == 1 else 10)
        nt.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    nt.column_dimensions["A"].width = 110
    wb.move_sheet("Synthèse", offset=-1)
    wb.save(out_path)


# --------------------------------------------------------------------------- main
def main():
    p = argparse.ArgumentParser(description="Étude de temps assistée par vidéo (vidéo -> Excel + Yamazumi)")
    p.add_argument("videos", nargs="+", type=Path, help="Une vidéo par poste")
    p.add_argument("--stations", nargs="+", help="Noms des postes (défaut : nom du fichier)")
    p.add_argument("--takt", type=float, help="Takt time en secondes (ligne sur le Yamazumi)")
    p.add_argument("--fps", type=float, default=1.0, help="Images analysées par seconde (moteur azure, défaut 1)")
    p.add_argument("--chunk-seconds", type=float, default=20.0, help="Durée d'un lot envoyé au modèle (azure)")
    p.add_argument("--width", type=int, default=768, help="Largeur max des images envoyées (azure)")
    p.add_argument("--start", type=float, default=0.0, help="Début de la fenêtre analysée (s)")
    p.add_argument("--end", type=float, default=None, help="Fin de la fenêtre analysée (s) - isoler un cycle")
    p.add_argument("--detail", choices=["low", "high", "auto"], default="auto")
    p.add_argument("--out", type=Path, default=None, help="Fichier Excel de sortie")
    p.add_argument("--backend", choices=["azure", "motion", "mock"], default="azure",
                   help="azure = modèle de vision ; motion = détection locale du mouvement (sans IA) ; mock = test")
    p.add_argument("--mock", action="store_true", help="Raccourci de --backend mock (données factices)")
    g = p.add_argument_group("options du moteur 'motion'")
    g.add_argument("--motion-fps", type=float, default=5.0, help="Images comparées par seconde (défaut 5)")
    g.add_argument("--motion-threshold", type=float, default=None,
                   help="Seuil d'activité en %% de pixels en mouvement (défaut : automatique)")
    g.add_argument("--min-phase", type=float, default=1.0, help="Durée minimale d'une phase en s (défaut 1.0)")
    g.add_argument("--smooth", type=float, default=0.6, help="Lissage du signal en s (défaut 0.6)")
    g.add_argument("--pixel-threshold", type=int, default=25,
                   help="Variation de luminosité (0-255) comptée comme mouvement (défaut 25)")
    g.add_argument("--blur", type=int, default=5, help="Flou gaussien impair contre le bruit caméra (défaut 5)")
    g.add_argument("--roi", default=None, help="Zone analysée en fractions de l'image : x,y,largeur,hauteur")
    args = p.parse_args()

    if args.stations and len(args.stations) != len(args.videos):
        sys.exit("--stations doit avoir autant de noms que de vidéos.")
    if args.blur % 2 == 0 or args.blur < 1:
        sys.exit("--blur doit être un entier impair (ex. 5).")
    stations = args.stations or [v.stem for v in args.videos]
    backend = "mock" if args.mock else args.backend
    model = AzureVision(detail=args.detail) if backend == "azure" else None

    results = {}
    for name, video in zip(stations, args.videos):
        if backend == "motion":
            steps, window = analyze_motion_video(video, args)
        else:
            steps, window = analyze_video(video, args, model)
        results[name] = (steps, window)
        out_json = video.with_suffix(".steps.json")
        out_json.write_text(json.dumps(steps, ensure_ascii=False, indent=2), encoding="utf-8")

    banners = {
        "mock": ("MODE TEST (--mock) : données factices, aucune analyse réelle de la vidéo.", "C00000"),
        "motion": ("MODE MOUVEMENT (sans IA) : phases d'activité / d'arrêt détectées par différence d'images, "
                   "sans interprétation. Saisir descriptions et catégories (VA / NVA / Gaspillage) dans l'onglet Étapes.",
                   "E65100"),
    }
    notes = {"motion": "Méthode : différence d'images successives (OpenCV, local, sans IA) ; seule la présence de "
                       "mouvement est mesurée, pas sa nature. Réglages : --motion-threshold, --roi, --min-phase."}
    out = args.out or Path(f"etude_temps_{datetime.now():%Y%m%d_%H%M}.xlsx")
    build_workbook(results, args.takt, out, banners.get(backend), notes.get(backend))
    print(f"\nTerminé : {out}")


if __name__ == "__main__":
    main()
